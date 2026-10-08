import openpyxl

ORIG = r'CN-PRJECT  DETSILS.xlsx'
EXPORT = r'Academic_Records_Updated.xlsx'

wb_orig_f = openpyxl.load_workbook(ORIG, data_only=False)
wb_exp_f = openpyxl.load_workbook(EXPORT, data_only=False)

ws_o = wb_orig_f['20-24']
ws_e = wb_exp_f['20-24']

print("--- 1. Checking Col 5 and Col 6 in Original vs Exported for Row 7 ---")
print("Original:")
print("  Cell(7, 5):", repr(ws_o.cell(7, 5).value), "fill:", ws_o.cell(7, 5).fill.fgColor.rgb if ws_o.cell(7, 5).fill and ws_o.cell(7, 5).fill.fgColor else None)
print("  Cell(7, 6):", repr(ws_o.cell(7, 6).value), "fill:", ws_o.cell(7, 6).fill.fgColor.rgb if ws_o.cell(7, 6).fill and ws_o.cell(7, 6).fill.fgColor else None)
print("Exported:")
print("  Cell(7, 5):", repr(ws_e.cell(7, 5).value), "fill:", ws_e.cell(7, 5).fill.fgColor.rgb if ws_e.cell(7, 5).fill and ws_e.cell(7, 5).fill.fgColor else None)
print("  Cell(7, 6):", repr(ws_e.cell(7, 6).value), "fill:", ws_e.cell(7, 6).fill.fgColor.rgb if ws_e.cell(7, 6).fill and ws_e.cell(7, 6).fill.fgColor else None)

print("\n--- 2. Checking #REF! in Original File ---")
orig_refs = []
for r in range(1, min(ws_o.max_row + 1, 135)):
    for c in range(1, min(ws_o.max_column + 1, 180)):
        v = ws_o.cell(r, c).value
        if v and isinstance(v, str) and '#REF!' in v:
            orig_refs.append((r, c, v))
print(f"Original file has {len(orig_refs)} #REF! formulas:")
for r, c, v in orig_refs[:5]:
    print(f"  Row {r}, Col {c}: {v}")

print("\n--- 3. Checking Row 7 Col 71 and Col 119 in DB and Arrears ---")
import json
with open(r'server\data\database.json', encoding='utf-8') as f:
    db = json.load(f)

# Find student at row 7: register 113020104001
s7 = [s for s in db['students'] if s['register_number'] == '113020104001'][0]
print(f"Student: {s7['id']} {s7['name']} {s7['register_number']}")
s7_arrears = [a for a in db['arrears'] if a['student_id'] == s7['id']]
print(f"Student arrears in DB ({len(s7_arrears)}):")
for a in s7_arrears:
    print(f"  {a['subject_code']} orig={a['original_grade']} cleared={a.get('cleared_grade')} sem={a.get('cleared_semester')} status={a['status']}")

print("\n--- 4. Checking original colors across all students on Col 71 (Sem 4) and Col 119 (Sem 6) ---")
for r in [7, 8, 9, 10, 11]:
    o_c4 = ws_o.cell(r, 71).fill.fgColor.rgb if ws_o.cell(r, 71).fill and ws_o.cell(r, 71).fill.fgColor else None
    e_c4 = ws_e.cell(r, 71).fill.fgColor.rgb if ws_e.cell(r, 71).fill and ws_e.cell(r, 71).fill.fgColor else None
    o_c6 = ws_o.cell(r, 119).fill.fgColor.rgb if ws_o.cell(r, 119).fill and ws_o.cell(r, 119).fill.fgColor else None
    e_c6 = ws_e.cell(r, 119).fill.fgColor.rgb if ws_e.cell(r, 119).fill and ws_e.cell(r, 119).fill.fgColor else None
    print(f"Row {r}: Sem 4 Orig={o_c4} Exp={e_c4} | Sem 6 Orig={o_c6} Exp={e_c6}")
