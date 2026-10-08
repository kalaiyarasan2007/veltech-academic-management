import openpyxl
import json

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws_main = wb['20-24']

# 1. Subjects per semester
# We mapped the grade columns for each semester:
sem_subject_cols = {
    1: [
        {'code': 'HS8151', 'title': 'Communicative English', 'credit': 4, 'col': 5},
        {'code': 'MA8151', 'title': 'Engineering Maths I', 'credit': 4, 'col': 7},
        {'code': 'PH8151', 'title': 'Engineering Physics', 'credit': 3, 'col': 9},
        {'code': 'CY8151', 'title': 'Engineering Chemistry', 'credit': 3, 'col': 11},
        {'code': 'GE8151', 'title': 'Python Pgmming', 'credit': 3, 'col': 13},
        {'code': 'GE8152', 'title': 'Engineering Graphics', 'credit': 4, 'col': 15},
        {'code': 'GE8161', 'title': 'Python Pgmming Laboratory', 'credit': 2, 'col': 17},
        {'code': 'BS8161', 'title': 'Physics and Chemistry Laboratory', 'credit': 2, 'col': 19},
    ],
    2: [
        {'code': 'HS8251', 'title': 'Technical English', 'credit': 4, 'col': 25},
        {'code': 'MA8251', 'title': 'Engineering Mathmatics II', 'credit': 4, 'col': 27},
        {'code': 'PH 8252', 'title': 'Physics and Information Science', 'credit': 3, 'col': 29},
        {'code': 'BE8255', 'title': 'BEEE', 'credit': 3, 'col': 31},
        {'code': 'GE8291', 'title': 'EVS', 'credit': 3, 'col': 33},
        {'code': 'CS8251', 'title': 'Programming in c', 'credit': 3, 'col': 35},
        {'code': 'GE8261', 'title': 'Engineering Practices Lab', 'credit': 2, 'col': 37},
        {'code': 'CS8261', 'title': 'C Pgmming Lab', 'credit': 2, 'col': 39},
    ],
    3: [
        {'code': 'MA8351', 'title': 'Discrete Mathematics', 'credit': 4, 'col': 47},
        {'code': 'CS8351', 'title': 'Digital Principles and System Design', 'credit': 4, 'col': 49},
        {'code': 'CS8391', 'title': 'Data Structure', 'credit': 3, 'col': 51},
        {'code': 'CS8392', 'title': 'OBJECT ORIENTEDPROGRAMMING', 'credit': 3, 'col': 53},
        {'code': 'EC8395', 'title': 'COMMUNICATION ENGINEERING', 'credit': 3, 'col': 55},
        {'code': 'CS8381', 'title': 'DATA STRUCTURES LABORATORY', 'credit': 2, 'col': 57},
        {'code': 'CS8383', 'title': 'OBJECT ORIENTED PROGRAMMING LABORATORY', 'credit': 2, 'col': 59},
        {'code': 'CS8382', 'title': 'DIGITAL SYSTEMS LABORATORY', 'credit': 2, 'col': 61},
        {'code': 'HS8381', 'title': 'INTERPERSONAL SKILLS', 'credit': 1, 'col': 63},
    ],
    4: [
        {'code': 'MA8402', 'title': 'PROBABILITY AND QUEUEING THEORY', 'credit': 4, 'col': 71},
        {'code': 'CS8491', 'title': 'COMPUTER ARCHITECTURE', 'credit': 3, 'col': 73},
        {'code': 'CS8492', 'title': 'DATABASE MANAGEMENT SYSTEMS', 'credit': 3, 'col': 75},
        {'code': 'CS8451', 'title': 'DESIGN AND ANALYSIS OF ALGORITHMS', 'credit': 3, 'col': 77},
        {'code': 'CS8493', 'title': 'OPERATING SYSTEMS', 'credit': 3, 'col': 79},
        {'code': 'CS8494', 'title': 'SOFTWARE ENGINEERING', 'credit': 3, 'col': 81},
        {'code': 'CS8481', 'title': 'DATABASE MANAGEMENT SYSTEMS LABORATORY', 'credit': 2, 'col': 83},
        {'code': 'CS8461', 'title': 'OPERATING SYSTEMS LABORATORY', 'credit': 2, 'col': 85},
        {'code': 'HS8461', 'title': 'ADVANCED READING AND WRITING', 'credit': 1, 'col': 87},
    ],
    5: [
        {'code': 'MA8551', 'title': 'ALGEBRA AND NUMBER THEORY', 'credit': 4, 'col': 95},
        {'code': 'CS8591', 'title': 'COMPUTER NETWORK', 'credit': 3, 'col': 97},
        {'code': 'EC8691', 'title': 'MICROPROCESSORS AND MICROCONTROLLERS', 'credit': 3, 'col': 99},
        {'code': 'CS8501', 'title': 'THEORY OF COMPUTATION', 'credit': 3, 'col': 101},
        {'code': 'CS8592', 'title': 'OBJECT ORIENTED ANALYSIS AND DESIGN', 'credit': 3, 'col': 103},
        {'code': 'OMF551', 'title': 'PRODUCT DESIGN AND DEVELOPMENT', 'credit': 3, 'col': 105},
        {'code': 'EC8681', 'title': 'MICROPROCESSORS AND MICROCONTROLLERS LABORATORY', 'credit': 2, 'col': 107},
        {'code': 'CS8582', 'title': 'OBJECT ORIENTED ANALYSIS AND DESIGN LABORATORY', 'credit': 2, 'col': 109},
        {'code': 'CS8581', 'title': 'NETWORKS LABORATORY', 'credit': 2, 'col': 111},
    ],
    6: [
        {'code': 'CS8651', 'title': 'INTERNET PROGRAMMING', 'credit': 3, 'col': 119},
        {'code': 'CS8691', 'title': 'ARTIFICIAL INTELLIGENCE', 'credit': 3, 'col': 121},
        {'code': 'CS8601', 'title': 'MOBILE COMPUTING', 'credit': 3, 'col': 123},
        {'code': 'CS8602', 'title': 'COMPILER DESIGN', 'credit': 4, 'col': 125},
        {'code': 'CS8603', 'title': 'DISTRIBUTED SYSTEMS', 'credit': 3, 'col': 127},
        {'code': 'IT8076', 'title': 'SOFTWARE TESTING', 'credit': 3, 'col': 129},
        {'code': 'CS8661', 'title': 'INTERNET PROGRAMMING LABORATORY', 'credit': 2, 'col': 131},
        {'code': 'CS8662', 'title': 'MOBILE APPLICATION DEVELOPMENT LABORATORY', 'credit': 2, 'col': 133},
        {'code': 'CS8611', 'title': 'MINI PROJECT', 'credit': 1, 'col': 135},
        {'code': 'HS8581', 'title': 'PROFESSIONAL COMMUNICATION', 'credit': 1, 'col': 137},
    ],
    7: [
        {'code': 'MG8591', 'title': 'PRINCIPLES OF MANAGEMENT', 'credit': 3, 'col': 145},
        {'code': 'CS8792', 'title': 'CRYPTOGRAPHY AND NETWORK SECURITY', 'credit': 3, 'col': 147},
        {'code': 'CS8791', 'title': 'CLOUD COMPUTING', 'credit': 3, 'col': 149},
        {'code': 'OME753', 'title': 'SYSTEMS ENGINEERING', 'credit': 3, 'col': 151},
        {'code': 'IT8075', 'title': 'SOFTWARE PROJECT MANAGEMENT', 'credit': 3, 'col': 153},
        {'code': 'CS8079', 'title': 'HUMAN COMPUTER INTERACTION', 'credit': 3, 'col': 155},
        {'code': 'CS8711', 'title': 'CLOUD COMPUTING LABORATORY', 'credit': 2, 'col': 157},
        {'code': 'IT8761', 'title': 'SECURITY LABORATORY', 'credit': 2, 'col': 159},
    ],
    8: [
        {'code': 'IT8073', 'title': 'INFORMATION SECURITY', 'credit': 3, 'col': 167},
        {'code': 'CS8080', 'title': 'INFORMATION RETRIEVAL TECHNIQUES', 'credit': 3, 'col': 169},
        {'code': 'CS8811', 'title': 'PROJECT WORK', 'credit': 10, 'col': 171},
    ]
}

sem_colors = {
    1: '#92D050',
    2: '#B3CEFA',
    3: '#FFA766',
    4: '#00B0F0',
    5: '#FFFF00',
    6: '#F28E85',
    7: '#CC9900',
    8: '#E2E8F0'
}

color_to_sem = {
    'FF92D050': 1,
    'FFB3CEFA': 2,
    'FFFFA766': 3,
    'FF00B0F0': 4,
    'FFFFFF00': 5,
    'FFF28E85': 6,
    'FFCC9900': 7,
}

students_data = []

for r in range(7, 133):
    sno = ws_main.cell(r, 2).value
    reg = ws_main.cell(r, 3).value
    name = ws_main.cell(r, 4).value
    
    reg_str = str(reg).strip() if reg else ""
    name_str = str(name).strip() if name else ""
    sno_str = str(sno).strip() if sno else ""
    
    if not reg_str or not reg_str.isdigit() or len(reg_str) < 8:
        continue
    
    student = {
        'excel_row': r,
        'sno': sno_str,
        'register_number': reg_str,
        'name': name_str,
        'roll_number': f"20CS{int(sno_str):03d}" if sno_str.isdigit() else f"20CS{r:03d}",
        'batch': '2020-2024',
        'department': 'Department of Computer Science and Engineering',
        'semesters': {}
    }
    
    for sem in range(1, 9):
        sem_subjects = []
        for s in sem_subject_cols[sem]:
            col_num = s['col']
            cell = ws_main.cell(r, col_num)
            grade_val = str(cell.value).strip() if cell.value is not None else ""
            
            # Check cell fill color
            cell_fill = cell.fill
            fg = cell_fill.fgColor if cell_fill else None
            rgb = fg.rgb if (fg and fg.type == 'rgb') else None
            
            # Arrear clearance detection
            cleared_in_sem = None
            if rgb in color_to_sem and color_to_sem[rgb] != sem:
                cleared_in_sem = color_to_sem[rgb]
            
            # Determine grade point
            gp_map = {'O': 10, 'A+': 9, 'A': 8, 'B+': 7, 'B': 6, 'C': 5}
            grade_point = gp_map.get(grade_val, 0)
            
            sem_subjects.append({
                'subject_code': s['code'],
                'subject_name': s['title'],
                'credits': s['credit'],
                'grade': grade_val,
                'grade_point': grade_point,
                'fill_rgb': rgb,
                'cleared_in_sem': cleared_in_sem
            })
        
        student['semesters'][str(sem)] = sem_subjects
    
    students_data.append(student)

print(f"Parsed {len(students_data)} valid students from sheet '20-24'.")

with open('extracted_students.json', 'w') as f:
    json.dump(students_data, f, indent=2)

print("Saved extracted_students.json successfully!")
