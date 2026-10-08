import urllib.request
import openpyxl
import io

url = 'http://localhost:5000/api/excel/export'
req = urllib.request.Request(url)
with urllib.request.urlopen(req) as resp:
    excel_bytes = resp.read()

wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
ws = wb['20-24']

print("=== INSPECTING EXPORTED EXCEL ROW 7 FOR COLS 47 TO 64 ===")
for c in range(47, 65):
    c_let = openpyxl.utils.get_column_letter(c)
    val4 = ws.cell(4, c).value
    val5 = ws.cell(5, c).value
    val7 = ws.cell(7, c).value
    print(f"Col {c:2d} ({c_let:2s}): R4='{val4}' | R5='{val5}' | R7='{val7}'")
