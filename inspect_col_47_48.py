import openpyxl

wb_formula = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=False)
wb_value = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)

ws_f = wb_formula['20-24']
ws_v = wb_value['20-24']

print("=== INSPECTING COLS 47 TO 54 (SEM 3) ===")
for c in range(47, 55):
    c_let = openpyxl.utils.get_column_letter(c)
    val_f = ws_f.cell(7, c).value
    val_v = ws_v.cell(7, c).value
    r4_f = ws_f.cell(4, c).value
    print(f"Col {c:2d} ({c_let:2s}): R4='{r4_f}' | R7 Formula='{val_f}' | R7 Value='{val_v}'")
