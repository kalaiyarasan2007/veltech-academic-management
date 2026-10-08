import openpyxl
import json

wb_formula = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=False)
ws = wb_formula['20-24']

# 1. Dynamically scan row 4 to find all subject columns in master template
subjects_list = []
current_sem = 1

for c in range(5, ws.max_column + 1):
    val4 = ws.cell(4, c).value
    val5 = ws.cell(5, c).value
    val6 = ws.cell(6, c).value

    str4 = str(val4).strip() if val4 else ""
    str5 = str(val5).strip() if val5 else ""
    str6 = str(val6).strip() if val6 else ""

    # Detect semester boundary from arrear columns
    if "Arrear" in str4 or "Arrear" in str5:
        if "1st" in str4 or "1st" in str5 or "Ist" in str4:
            current_sem = 2
        elif "2nd" in str4 or "Ist year" in str4:
            current_sem = 3
        elif "3rd" in str4:
            current_sem = 4
        elif "4th" in str4:
            current_sem = 5
        elif "5th" in str4:
            current_sem = 6
        elif "6th" in str4:
            current_sem = 7
        elif "7th" in str4:
            current_sem = 8

    # If row 4 has a subject code and row 6 has credit value
    if str4 and str6.isdigit() and not str4.startswith("Total") and not str4.startswith("GPA") and not str4.startswith("CGPA") and not str4.startswith("Arrear"):
        subjects_list.append({
            'semester_number': current_sem,
            'subject_code': str4,
            'subject_name': str5,
            'credits': int(str6),
            'grade_col': c,
            'gp_col': c + 1
        })

print(f"Dynamically mapped {len(subjects_list)} subjects across all 8 semesters from master template.")
for s in subjects_list:
    c_let = openpyxl.utils.get_column_letter(s['grade_col'])
    print(f"  Sem {s['semester_number']} | Col {s['grade_col']:3d} ({c_let:3s}): {s['subject_code']:8s} - {s['subject_name']:35s} - Credits: {s['credits']}")

color_to_sem = {
    'FF92D050': 1,
    'FFB3CEFA': 2,
    'FFFFA766': 3,
    'FF00B0F0': 4,
    'FFFFFF00': 5,
    'FFF28E85': 6,
    'FFCC9900': 7,
}

gp_map = {'O': 10, 'A+': 9, 'A': 8, 'B+': 7, 'B': 6, 'C': 5}

students_data = []

for r in range(7, 133):
    sno = ws.cell(r, 2).value
    reg = ws.cell(r, 3).value
    name = ws.cell(r, 4).value

    reg_str = str(reg).strip() if reg else ""
    name_str = str(name).strip() if name else ""
    sno_str = str(sno).strip() if sno else ""

    if not reg_str or not reg_str.isdigit() or len(reg_str) < 8:
        continue

    student = {
        'excel_row': r,
        'sno': sno_str,
        'register_number': reg_str,
        'name': name_str,
        'roll_number': f"20CS{int(sno_str):03d}" if sno_str.isdigit() else f"20CS{r:03d}",
        'batch': '2020-2024',
        'department': 'Department of Computer Science and Engineering',
        'semesters': {}
    }

    for sem in range(1, 9):
        sem_subs = [s for s in subjects_list if s['semester_number'] == sem]
        sub_records = []

        for sub in sem_subs:
            cell = ws.cell(r, sub['grade_col'])
            grade_val = str(cell.value).strip() if cell.value is not None else ""

            # Check fill color for arrear clearance detection
            cell_fill = cell.fill
            fg = cell_fill.fgColor if cell_fill else None
            rgb = fg.rgb if (fg and fg.type == 'rgb') else None

            cleared_in_sem = None
            if rgb in color_to_sem and color_to_sem[rgb] != sem:
                cleared_in_sem = color_to_sem[rgb]

            grade_point = gp_map.get(grade_val, 0)

            sub_records.append({
                'subject_code': sub['subject_code'],
                'subject_name': sub['subject_name'],
                'credits': sub['credits'],
                'grade': grade_val,
                'grade_point': grade_point,
                'fill_rgb': rgb,
                'cleared_in_sem': cleared_in_sem
            })

        student['semesters'][str(sem)] = sub_records

    students_data.append(student)

print(f"\nParsed {len(students_data)} valid students.")

with open('extracted_students.json', 'w') as f:
    json.dump(students_data, f, indent=2)

print("Saved updated extracted_students.json successfully!")
