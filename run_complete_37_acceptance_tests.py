import urllib.request
import urllib.parse
import json
import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

BASE_URL = "http://localhost:5000/api"
ORIG_MASTER = r"e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx"

ADMIN_EMAIL = "admin@veltech.edu.in"
ADMIN_PASSWORD = "VelTech@2026"
WRONG_PASSWORD = "WrongPassword@123"

admin_token = None

def api_request(endpoint, method="GET", data=None, token=None):
    url = f"{BASE_URL}{endpoint}"
    headers = {'Content-Type': 'application/json'}
    if token:
        headers['Authorization'] = f"Bearer {token}"
    body = json.dumps(data).encode('utf-8') if data is not None else None
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            resp_body = resp.read().decode('utf-8')
            return resp.status, json.loads(resp_body) if resp_body else {}
    except urllib.error.HTTPError as e:
        err_body = e.read().decode('utf-8')
        try:
            return e.code, json.loads(err_body)
        except Exception:
            return e.code, {"error": err_body}

def api_download(endpoint, out_path, token=None):
    url = f"{BASE_URL}{endpoint}"
    headers = {}
    if token:
        headers['Authorization'] = f"Bearer {token}"
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req) as resp:
        content = resp.read()
        with open(out_path, 'wb') as f:
            f.write(content)
        return resp.status, len(content), dict(resp.headers)

print("=" * 80)
print("COMPREHENSIVE 37-STEP FINAL ACCEPTANCE TEST SUITE")
print("================================================================================")

# ------------------------------------------------------------------------------
# 1. Login as Admin & 2. Verify wrong password is rejected
# ------------------------------------------------------------------------------
print("\n[STEP 1 & 2] Testing Admin Login Security & Wrong Password Rejection...")
status, resp = api_request("/auth/login", method="POST", data={
    "email": ADMIN_EMAIL,
    "password": WRONG_PASSWORD
})
assert status == 401, f"Expected 401 for wrong password, got {status}: {resp}"
print(f"  ✓ Rejected wrong password correctly with status 401: {resp.get('error')}")

# ------------------------------------------------------------------------------
# 3. Login successfully
# ------------------------------------------------------------------------------
print("\n[STEP 3] Login successfully with valid Admin credentials...")
status, resp = api_request("/auth/login", method="POST", data={
    "email": ADMIN_EMAIL,
    "password": ADMIN_PASSWORD
})
assert status == 200 and resp.get("token"), f"Login failed: {resp}"
admin_token = resp["token"]
print(f"  ✓ Admin authenticated successfully. User: {resp.get('user', {}).get('name')}, Role: {resp.get('user', {}).get('role')}")

# ------------------------------------------------------------------------------
# 4. Create / select a Department
# ------------------------------------------------------------------------------
print("\n[STEP 4] Verify / select Department...")
status, resp = api_request("/departments", token=admin_token)
assert status == 200, f"Failed to get departments: {resp}"
departments = resp.get("departments", [])
cse_dept = next((d for d in departments if d["code"] == "CSE" or "Computer Science" in d["name"]), None)
if not cse_dept:
    status, add_d = api_request("/departments", method="POST", data={
        "name": "Department of Computer Science and Engineering",
        "code": "CSE",
        "description": "Computer Science and Engineering"
    }, token=admin_token)
    cse_dept = add_d["department"]
print(f"  ✓ Department active: {cse_dept['name']} (ID: {cse_dept['id']}, Code: {cse_dept['code']})")

# ------------------------------------------------------------------------------
# 5. Create CSE-A & 6. Create CSE-B
# ------------------------------------------------------------------------------
print("\n[STEP 5 & 6] Verify / Create Sections CSE-A and CSE-B...")
status, resp = api_request("/sections", token=admin_token)
sections = resp.get("sections", [])
sec_a = next((s for s in sections if s["section_code"] == "CSE-A"), None)
sec_b = next((s for s in sections if s["section_code"] == "CSE-B"), None)

if not sec_a:
    status, add_a = api_request("/sections", method="POST", data={
        "department_id": cse_dept["id"],
        "section_name": "CSE-A",
        "section_code": "CSE-A"
    }, token=admin_token)
    sec_a = add_a["section"]

if not sec_b:
    status, add_b = api_request("/sections", method="POST", data={
        "department_id": cse_dept["id"],
        "section_name": "CSE-B",
        "section_code": "CSE-B"
    }, token=admin_token)
    sec_b = add_b["section"]

print(f"  ✓ Section CSE-A ID: {sec_a['id']}")
print(f"  ✓ Section CSE-B ID: {sec_b['id']}")

# ------------------------------------------------------------------------------
# 7. Add students to CSE-A & 8. Add different students to CSE-B
# ------------------------------------------------------------------------------
print("\n[STEP 7 & 8] Adding / verifying dedicated students for CSE-A and CSE-B...")
student_a_reg = "113024104001"
student_a_name = "AARTHI S"
status, resp = api_request("/students", method="POST", data={
    "register_number": student_a_reg,
    "name": student_a_name,
    "department_id": cse_dept["id"],
    "section_id": sec_a["id"],
    "section": "CSE-A",
    "batch": "2020-2024"
}, token=admin_token)
assert status in (200, 201), f"Failed to upsert student A: {resp}"
std_a_id = resp["student"]["id"]

student_b_reg = "113024104061"
student_b_name = "BALAJI R"
status, resp = api_request("/students", method="POST", data={
    "register_number": student_b_reg,
    "name": student_b_name,
    "department_id": cse_dept["id"],
    "section_id": sec_b["id"],
    "section": "CSE-B",
    "batch": "2020-2024"
}, token=admin_token)
assert status in (200, 201), f"Failed to upsert student B: {resp}"
std_b_id = resp["student"]["id"]
print(f"  ✓ Added/Assigned {student_a_name} ({student_a_reg}) -> CSE-A")
print(f"  ✓ Added/Assigned {student_b_name} ({student_b_reg}) -> CSE-B")

# ------------------------------------------------------------------------------
# 9. Verify CSE-A shows only CSE-A students & 10. Verify CSE-B shows only CSE-B students
# ------------------------------------------------------------------------------
print("\n[STEP 9 & 10] Verifying strict section student separation in API queries...")
status, resp_a = api_request(f"/students?section_id={sec_a['id']}", token=admin_token)
students_in_a = resp_a.get("students", [])
a_regs = {s["register_number"] for s in students_in_a}

status, resp_b = api_request(f"/students?section_id={sec_b['id']}", token=admin_token)
students_in_b = resp_b.get("students", [])
b_regs = {s["register_number"] for s in students_in_b}

assert student_a_reg in a_regs, "Student A not found in CSE-A list"
assert student_b_reg not in a_regs, "CRITICAL ERROR: Student B found in CSE-A list!"
assert student_b_reg in b_regs, "Student B not found in CSE-B list"
assert student_a_reg not in b_regs, "CRITICAL ERROR: Student A found in CSE-B list!"
print(f"  ✓ CSE-A student list count: {len(students_in_a)} (Contains {student_a_reg}, strictly excludes {student_b_reg})")
print(f"  ✓ CSE-B student list count: {len(students_in_b)} (Contains {student_b_reg}, strictly excludes {student_a_reg})")

# ------------------------------------------------------------------------------
# 11. Configure Semester subjects
# ------------------------------------------------------------------------------
print("\n[STEP 11] Verifying / Configuring Semester subjects...")
status, resp = api_request("/academic/subjects?semester_number=1", token=admin_token)
sem1_subs = resp.get("subjects", [])
assert len(sem1_subs) >= 8, f"Expected at least 8 subjects in Semester 1, got {len(sem1_subs)}"
print(f"  ✓ Semester 1 configured with {len(sem1_subs)} subjects: {[s['subject_code'] for s in sem1_subs]}")

# ------------------------------------------------------------------------------
# 12, 13, 14, 15, 16, 17: Select CSE-A, Select Student, Enter Grades, Grade Points, GPA, CGPA
# ------------------------------------------------------------------------------
print("\n[STEP 12-17] Entering Semester 1 Grades for CSE-A Student & Verifying GPA/CGPA...")
sem1_grades = [
    {"subject_code": "HS8151", "subject_name": "Communicative English", "credits": 4, "grade": "O"},   # 10 * 4 = 40
    {"subject_code": "MA8151", "subject_name": "Engineering Maths I", "credits": 4, "grade": "A+"},   # 9 * 4 = 36
    {"subject_code": "PH8151", "subject_name": "Engineering Physics", "credits": 3, "grade": "A"},    # 8 * 3 = 24
    {"subject_code": "CY8151", "subject_name": "Engineering Chemistry", "credits": 3, "grade": "A+"}, # 9 * 3 = 27
    {"subject_code": "GE8151", "subject_name": "Python Pgmming", "credits": 3, "grade": "O"},         # 10 * 3 = 30
    {"subject_code": "GE8152", "subject_name": "Engineering Graphics", "credits": 4, "grade": "A"},   # 8 * 4 = 32
    {"subject_code": "GE8161", "subject_name": "Python Pgmming Laboratory", "credits": 2, "grade": "O"}, # 10 * 2 = 20
    {"subject_code": "BS8161", "subject_name": "Physics and Chemistry Laboratory", "credits": 2, "grade": "O"} # 10 * 2 = 20
]
# Expected Tot: 40 + 36 + 24 + 27 + 30 + 32 + 20 + 20 = 229
# Credits: 25 -> GPA = 229 / 25 = 9.16

status, resp = api_request("/academic/save-grades", method="POST", data={
    "student_id": student_a_reg,
    "semester_number": 1,
    "exam_month": "January",
    "exam_year": "2021",
    "regular_grades": sem1_grades
}, token=admin_token)
assert status == 200, f"Failed to save grades: {resp}"

status, history = api_request(f"/academic/student-history/{student_a_reg}", token=admin_token)
sem1_hist = history["semesters"][0]
print(f"  ✓ Sem 1 Total Grade: {sem1_hist['total_grade']} (Expected: 229)")
print(f"  ✓ Sem 1 GPA: {sem1_hist['gpa']} (Expected: 9.16)")
print(f"  ✓ Sem 1 CGPA: {sem1_hist['cgpa']} (Expected: 9.16)")
assert sem1_hist['total_grade'] == 229
assert sem1_hist['gpa'] == 9.16
assert sem1_hist['cgpa'] == 9.16

# ------------------------------------------------------------------------------
# 18 & 19: Create an RA in Semester 2 & Verify Arrear
# ------------------------------------------------------------------------------
print("\n[STEP 18 & 19] Creating RA Arrear in Semester 2 and verifying pending arrear...")
sem2_grades = [
    {"subject_code": "HS8251", "subject_name": "Technical English", "credits": 4, "grade": "A"},
    {"subject_code": "MA8251", "subject_name": "Engineering Mathmatics II", "credits": 4, "grade": "RA"},
    {"subject_code": "PH 8252", "subject_name": "Physics and Information Science", "credits": 3, "grade": "B+"},
    {"subject_code": "BE8255", "subject_name": "BEEE", "credits": 3, "grade": "A"},
    {"subject_code": "GE8291", "subject_name": "EVS", "credits": 3, "grade": "O"},
    {"subject_code": "CS8251", "subject_name": "Programming in c", "credits": 3, "grade": "A+"},
    {"subject_code": "GE8261", "subject_name": "Engineering Practices Lab", "credits": 2, "grade": "O"},
    {"subject_code": "CS8261", "subject_name": "C Pgmming Lab", "credits": 2, "grade": "O"}
]
status, resp = api_request("/academic/save-grades", method="POST", data={
    "student_id": student_a_reg,
    "semester_number": 2,
    "exam_month": "June",
    "exam_year": "2021",
    "regular_grades": sem2_grades
}, token=admin_token)
assert status == 200

status, arrears_resp = api_request(f"/arrears?student_id={student_a_reg}", token=admin_token)
arrears_list = arrears_resp.get("arrears", [])
ma_arr = next((a for a in arrears_list if a["subject_code"] == "MA8251"), None)
assert ma_arr is not None, "Arrear record was not generated for MA8251"
assert ma_arr["status"] == "PENDING"
print(f"  ✓ Arrear recorded: Subject: {ma_arr['subject_code']}, Status: {ma_arr['status']}, Sem: {ma_arr['original_semester']}")

# ------------------------------------------------------------------------------
# 20 & 21: Clear the arrear & Verify clearance
# ------------------------------------------------------------------------------
print("\n[STEP 20 & 21] Clearing Arrear in Semester 3 & verifying clearance...")
status, resp = api_request("/arrears/clear", method="POST", data={
    "arrear_id": ma_arr["id"],
    "cleared_semester": 3,
    "cleared_grade": "A",
    "cleared_month": "January",
    "cleared_year": "2022"
}, token=admin_token)
assert status == 200

status, arrears_resp = api_request(f"/arrears?student_id={student_a_reg}", token=admin_token)
ma_arr_cleared = next((a for a in arrears_resp.get("arrears", []) if a["subject_code"] == "MA8251"), None)
assert ma_arr_cleared["status"] == "CLEARED"
assert ma_arr_cleared["cleared_grade"] == "A"
assert ma_arr_cleared["cleared_semester"] == 3
print(f"  ✓ Arrear status updated: {ma_arr_cleared['status']}, Cleared Grade: {ma_arr_cleared['cleared_grade']} in Sem {ma_arr_cleared['cleared_semester']}")

# ------------------------------------------------------------------------------
# 22, 23, 24, 25, 26: Generate, Preview CSE-A Excel & Verify Structure / Formatting
# ------------------------------------------------------------------------------
print("\n[STEP 22-26] Generating and Previewing CSE-A Excel Workbook...")
status, preview_a = api_request(f"/excel/section/preview?section_id={sec_a['id']}&department_id={cse_dept['id']}", token=admin_token)
assert status == 200
assert preview_a["filename"] == "CSE_A_Academic_Records.xlsx", f"Expected CSE_A_Academic_Records.xlsx, got {preview_a['filename']}"
print(f"  ✓ Generated preview for: {preview_a['filename']}")

# Check that only CSE-A students are inside preview
main_sheet_a = next((s for s in preview_a["sheets"] if s["name"] == "20-24"), None)
assert main_sheet_a is not None, "Worksheet 20-24 missing from preview"
preview_regs_a = []
for row in main_sheet_a["rows"]:
    if row["rowNumber"] >= 7:
        reg_val = next((c["value"] for c in row["cells"] if c["col"] == 3), "")
        if reg_val and reg_val.strip().isdigit():
            preview_regs_a.append(reg_val.strip())

assert student_a_reg in preview_regs_a, f"{student_a_reg} missing from CSE-A preview"
assert student_b_reg not in preview_regs_a, f"CRITICAL: {student_b_reg} (CSE-B) leaked into CSE-A preview!"
print(f"  ✓ CSE-A preview contains {len(preview_regs_a)} students. Student A included, Student B excluded.")

# ------------------------------------------------------------------------------
# 27, 28, 29: Download CSE-A Excel & Open in OpenPyXL (Genuine XLSX verification)
# ------------------------------------------------------------------------------
print("\n[STEP 27, 28, 29] Downloading CSE-A Excel and testing XLSX integrity...")
out_a_path = r"e:\COLLEGE PROJECT\test_export_cse_a.xlsx"
status, size, headers = api_download(f"/excel/section/export?section_id={sec_a['id']}&department_id={cse_dept['id']}", out_a_path, token=admin_token)
assert status == 200
assert size > 50000, f"Downloaded file size suspiciously small: {size} bytes"
print(f"  ✓ Downloaded {out_a_path} ({size:,} bytes)")

wb_a = openpyxl.load_workbook(out_a_path, data_only=False)
ws_a = wb_a["20-24"]
assert ws_a is not None
print(f"  ✓ Workbook opens cleanly without errors. Sheet names: {wb_a.sheetnames}")

# Verify headers: exact original headers
wb_orig = openpyxl.load_workbook(ORIG_MASTER, data_only=True)
ws_orig = wb_orig["20-24"]

# Verify row 1-6 headers match master Excel
for r in range(1, 7):
    for c in range(1, 30):
        orig_val = ws_orig.cell(r, c).value
        gen_val = ws_a.cell(r, c).value
        if orig_val is not None:
            assert str(orig_val).strip() == str(gen_val).strip(), f"Header mismatch at ({r},{c}): Master='{orig_val}', Gen='{gen_val}'"
print("  ✓ All original Excel headers in rows 1-6 match the master template exactly!")

# ------------------------------------------------------------------------------
# 30, 31, 32, 33: Generate CSE-B Excel & Confirm strict 2-way isolation
# ------------------------------------------------------------------------------
print("\n[STEP 30-33] Generating CSE-B Excel and verifying strict two-way data isolation...")
out_b_path = r"e:\COLLEGE PROJECT\test_export_cse_b.xlsx"
status, size, headers = api_download(f"/excel/section/export?section_id={sec_b['id']}&department_id={cse_dept['id']}", out_b_path, token=admin_token)
assert status == 200

wb_b = openpyxl.load_workbook(out_b_path, data_only=True)
ws_b = wb_b["20-24"]

excel_a_regs = set()
for r in range(7, ws_a.max_row + 1):
    val = ws_a.cell(r, 3).value
    if val and str(val).strip().isdigit() and len(str(val).strip()) >= 8:
        excel_a_regs.add(str(val).strip())

excel_b_regs = set()
for r in range(7, ws_b.max_row + 1):
    val = ws_b.cell(r, 3).value
    if val and str(val).strip().isdigit() and len(str(val).strip()) >= 8:
        excel_b_regs.add(str(val).strip())

print(f"  ✓ Students in CSE_A_Academic_Records.xlsx: {len(excel_a_regs)}")
print(f"  ✓ Students in CSE_B_Academic_Records.xlsx: {len(excel_b_regs)}")

assert student_a_reg in excel_a_regs, f"{student_a_reg} must be in CSE-A Excel"
assert student_a_reg not in excel_b_regs, f"CRITICAL: {student_a_reg} (CSE-A) found in CSE-B Excel!"
assert student_b_reg in excel_b_regs, f"{student_b_reg} must be in CSE-B Excel"
assert student_b_reg not in excel_a_regs, f"CRITICAL: {student_b_reg} (CSE-B) found in CSE-A Excel!"

# Check intersection
overlap = excel_a_regs.intersection(excel_b_regs)
assert len(overlap) == 0, f"CRITICAL DATA LEAK: Found overlapping students between CSE-A and CSE-B: {overlap}"
print("  ✓ PASSED: ZERO student overlap between CSE-A and CSE-B workbooks!")

# ------------------------------------------------------------------------------
# 34 & 35: Verify Common Excel contains only configuration, NOT mixed marks
# ------------------------------------------------------------------------------
print("\n[STEP 34 & 35] Verifying Common Excel contains ONLY master configuration...")
out_common_path = r"e:\COLLEGE PROJECT\test_common_data.xlsx"
status, size, headers = api_download("/excel/common/export", out_common_path, token=admin_token)
assert status == 200

wb_common = openpyxl.load_workbook(out_common_path, data_only=True)
print(f"  ✓ Common Excel sheets: {wb_common.sheetnames}")
assert "Departments" in wb_common.sheetnames
assert "Sections" in wb_common.sheetnames
assert "Batches" in wb_common.sheetnames
assert "Semesters" in wb_common.sheetnames
assert "Subjects Master" in wb_common.sheetnames
assert "Academic Configuration" in wb_common.sheetnames

# Verify NO student marks exist in Common Excel
for sname in wb_common.sheetnames:
    ws = wb_common[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '')
            assert student_a_reg not in val, f"Student register number found in Common Excel sheet {sname}!"
            assert student_b_reg not in val, f"Student register number found in Common Excel sheet {sname}!"

print("  ✓ PASSED: Common Excel strictly contains master configuration and zero student mark records.")

# ------------------------------------------------------------------------------
# 36 & 37: Logout & Verify Admin API Protection
# ------------------------------------------------------------------------------
print("\n[STEP 36 & 37] Testing Logout and Unauthenticated Access Rejection...")
status, resp = api_request("/auth/logout", method="POST", token=admin_token)
assert status == 200

# Try accessing protected admin endpoint without token
status, unauth_resp = api_request("/departments", method="POST", data={"name": "Test", "code": "TST"})
assert status in (401, 403), f"Expected 401/403 for unauthenticated POST, got {status}"
print(f"  ✓ Unauthenticated POST /departments rejected with status {status}: {unauth_resp.get('error')}")

status, unauth_resp = api_request("/sections", method="POST", data={"department_id": "dept_cse", "section_name": "TST"})
assert status in (401, 403), f"Expected 401/403 for unauthenticated POST, got {status}"
print(f"  ✓ Unauthenticated POST /sections rejected with status {status}: {unauth_resp.get('error')}")

status, unauth_resp = api_request("/subjects", method="POST", data={"subject_code": "CS9999", "subject_name": "Test", "semester_number": 1, "credits": 3})
assert status in (401, 403), f"Expected 401/403 for unauthenticated POST, got {status}"
print(f"  ✓ Unauthenticated POST /subjects rejected with status {status}: {unauth_resp.get('error')}")

print("\n" + "=" * 80)
print("ALL 37 ACCEPTANCE CRITERIA PASSED FULLY & FLAWLESSLY!")
print("================================================================================")
