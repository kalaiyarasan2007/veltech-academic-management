import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    students = []
    for r in range(7, ws.max_row + 1):
        sno = ws.cell(r, 2).value
        reg = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        if reg and str(reg).strip().isdigit() and len(str(reg).strip()) > 8:
            students.append((r, sno, str(reg).strip(), str(name).strip() if name else ""))
    print(f"Sheet '{sheetname}': {len(students)} students found (from row {students[0][0] if students else 'N/A'} to {students[-1][0] if students else 'N/A'})")

# Print first 5 and last 5 students in 20-24
ws = wb['20-24']
students_2024 = []
for r in range(7, ws.max_row + 1):
    reg = ws.cell(r, 3).value
    name = ws.cell(r, 4).value
    if reg and str(reg).strip().isdigit() and len(str(reg).strip()) > 8:
        students_2024.append((r, str(reg).strip(), str(name).strip() if name else ""))

print("\nFirst 5 students:")
for s in students_2024[:5]:
    print("  Row", s[0], ":", s[1], "-", s[2])

print("\nLast 5 students:")
for s in students_2024[-5:]:
    print("  Row", s[0], ":", s[1], "-", s[2])
