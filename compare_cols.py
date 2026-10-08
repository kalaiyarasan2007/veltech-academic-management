import openpyxl

ORIG   = r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx'
EXPORT = r'e:\COLLEGE PROJECT\test_export.xlsx'

wb_orig = openpyxl.load_workbook(ORIG, data_only=True)
wb_exp  = openpyxl.load_workbook(EXPORT, data_only=True)

ws_o = wb_orig['20-24']
ws_e = wb_exp['20-24']

print(f'Original max columns: {ws_o.max_column}')
print(f'Exported max columns: {ws_e.max_column}')

print()
print('ORIGINAL row 4 headers (cols 1-50):')
for c in range(1, 51):
    v = ws_o.cell(4, c).value
    if v: print(f'  col {c}: {repr(str(v).strip()[:30])}')

print()
print('EXPORTED row 4 headers (cols 1-50):')
for c in range(1, 51):
    v = ws_e.cell(4, c).value
    if v: print(f'  col {c}: {repr(str(v).strip()[:30])}')
