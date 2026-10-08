import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']

sem_colors = {
    1: 'FF92D050',
    2: 'FFB3CEFA',
    3: 'FFFFA766',
    4: 'FF00B0F0',
    5: 'FFFFFF00',
    6: 'FFF28E85',
    7: 'FFCC9900',
    8: '00000000' # No fill
}

sem_grade_cols = {
    1: [5, 7, 9, 11, 13, 15, 17, 19],
    2: [25, 27, 29, 31, 33, 35, 37, 39],
    3: [47, 49, 51, 53, 55, 57, 59, 61, 63],
    4: [71, 73, 75, 77, 79, 81, 83, 85, 87],
    5: [95, 97, 99, 101, 103, 105, 107, 109, 111],
    6: [119, 121, 123, 125, 127, 129, 131, 133, 135, 137],
    7: [145, 147, 149, 151, 153, 155, 157, 159],
    8: [167, 169, 171]
}

print("=== CHECKING FOR ARREARS & SPECIAL COLORS IN EXCEL ===")
for r in range(7, 133):
    reg = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value else ""
    name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value else ""
    if not reg or not reg.isdigit():
        continue
    
    special_cells = []
    for sem, cols in sem_grade_cols.items():
        expected_color = sem_colors[sem]
        for c in cols:
            val = str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else ""
            subj_code = str(ws.cell(4, c).value).strip()
            
            cell_fill = ws.cell(r, c).fill
            fg = cell_fill.fgColor if cell_fill else None
            actual_color = fg.rgb if (fg and fg.type == 'rgb') else 'NONE'
            
            # Check if grade is RA/U/AB/WH/WD or if color is different from expected sem color
            is_arrear_grade = val in ['RA', 'U', 'AB', 'WH', 'WD']
            is_diff_color = (actual_color != expected_color) and not (sem == 8 and actual_color == 'NONE')
            
            if is_arrear_grade or is_diff_color:
                special_cells.append({
                    'sem': sem, 'col': c, 'code': subj_code, 
                    'grade': val, 'expected_color': expected_color, 'actual_color': actual_color
                })
    
    if special_cells:
        print(f"\nStudent Row {r:3d}: {reg} - {name}")
        for item in special_cells:
            print(f"  Sem {item['sem']} | Col {item['col']:3d} ({openpyxl.utils.get_column_letter(item['col'])}) | Subj: {item['code']:8s} | Grade: {item['grade']:4s} | ActualColor: {item['actual_color']} | ExpectedColor: {item['expected_color']}")
