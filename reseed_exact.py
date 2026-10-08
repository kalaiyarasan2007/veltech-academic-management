import openpyxl
import json

wb_formula = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=False)
ws = wb_formula['20-24']

sem_col_ranges = {
    1: range(5, 20, 2),    # Cols 5, 7, 9, 11, 13, 15, 17, 19
    2: range(25, 40, 2),   # Cols 25, 27, 29, 31, 33, 35, 37, 39
    3: range(47, 64, 2),   # Cols 47, 49, 51, 53, 55, 57, 59, 61, 63
    4: range(71, 88, 2),   # Cols 71, 73, 75, 77, 79, 81, 83, 85, 87
    5: range(95, 112, 2),  # Cols 95, 97, 99, 101, 103, 105, 107, 109, 111
    6: range(119, 138, 2), # Cols 119, 121, 123, 125, 127, 129, 131, 133, 135, 137
    7: range(145, 160, 2), # Cols 145, 147, 149, 151, 153, 155, 157, 159
    8: range(167, 172, 2)  # Cols 167, 169, 171
}

subjects_list = []
for sem, col_range in sem_col_ranges.items():
    for c in col_range:
        val4 = ws.cell(4, c).value
        val5 = ws.cell(5, c).value
        val6 = ws.cell(6, c).value

        str4 = str(val4).strip() if val4 else ""
        str5 = str(val5).strip() if val5 else ""
        str6 = str(val6).strip() if val6 else ""

        if str4 and str6.isdigit():
            subjects_list.append({
                'semester_number': sem,
                'subject_code': str4,
                'subject_name': str5,
                'credits': int(str6),
                'grade_col': c,
                'gp_col': c + 1
            })

print(f"Mapped {len(subjects_list)} subjects across 8 semesters.")

color_to_sem = {
    'FF92D050': 1,
    'FFB3CEFA': 2,
    'FFFFA766': 3,
    'FF00B0F0': 4,
    'FFFFFF00': 5,
    'FFF28E85': 6,
    'FFCC9900': 7,
}

gp_map = {'O': 10, 'A+': 9, 'A': 8, 'B+': 7, 'B': 6, 'C': 5}

students_data = []

for r in range(7, 133):
    sno = ws.cell(r, 2).value
    reg = ws.cell(r, 3).value
    name = ws.cell(r, 4).value

    reg_str = str(reg).strip() if reg else ""
    name_str = str(name).strip() if name else ""
    sno_str = str(sno).strip() if sno else ""

    if not reg_str or not reg_str.isdigit() or len(reg_str) < 8:
        continue

    student = {
        'excel_row': r,
        'sno': sno_str,
        'register_number': reg_str,
        'name': name_str,
        'roll_number': f"20CS{int(sno_str):03d}" if sno_str.isdigit() else f"20CS{r:03d}",
        'batch': '2020-2024',
        'department': 'Department of Computer Science and Engineering',
        'semesters': {}
    }

    for sem in range(1, 9):
        sem_subs = [s for s in subjects_list if s['semester_number'] == sem]
        sub_records = []

        for sub in sem_subs:
            cell = ws.cell(r, sub['grade_col'])
            grade_val = str(cell.value).strip() if cell.value is not None else ""

            # Check fill color for arrear clearance detection
            cell_fill = cell.fill
            fg = cell_fill.fgColor if cell_fill else None
            rgb = fg.rgb if (fg and fg.type == 'rgb') else None

            cleared_in_sem = None
            if rgb in color_to_sem and color_to_sem[rgb] != sem:
                cleared_in_sem = color_to_sem[rgb]

            grade_point = gp_map.get(grade_val, 0)

            sub_records.append({
                'subject_code': sub['subject_code'],
                'subject_name': sub['subject_name'],
                'credits': sub['credits'],
                'grade': grade_val,
                'grade_point': grade_point,
                'fill_rgb': rgb,
                'cleared_in_sem': cleared_in_sem
            })

        student['semesters'][str(sem)] = sub_records

    students_data.append(student)

print(f"Parsed {len(students_data)} valid students.")

with open('extracted_students.json', 'w') as f:
    json.dump(students_data, f, indent=2)

print("Saved clean extracted_students.json successfully!")
