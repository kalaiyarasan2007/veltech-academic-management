import json
import openpyxl

with open(r'server\data\database.json', encoding='utf-8') as f:
    db = json.load(f)

wb = openpyxl.load_workbook(r'CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']
orig_regs = {}
for r in range(7, ws.max_row+1):
    reg = ws.cell(r, 3).value
    name = ws.cell(r, 4).value
    if reg and str(reg).strip().isdigit() and len(str(reg).strip()) >= 8:
        orig_regs[str(reg).strip()] = (r, str(name).strip() if name else '')

db_regs = {s['register_number']: s['name'] for s in db['students']}

print(f"Original Excel student count: {len(orig_regs)}")
print(f"Database student count: {len(db_regs)}")

in_db_not_orig = set(db_regs.keys()) - set(orig_regs.keys())
in_orig_not_db = set(orig_regs.keys()) - set(db_regs.keys())

print(f"In DB but NOT in original Excel: {len(in_db_not_orig)}")
for reg in sorted(in_db_not_orig):
    print(f"  Extra in DB -> Register: {reg}, Name: {db_regs[reg]}")

print(f"In original Excel but NOT in DB: {len(in_orig_not_db)}")
for reg in sorted(in_orig_not_db):
    print(f"  Missing from DB -> Register: {reg}, Name: {orig_regs[reg][1]} (Row {orig_regs[reg][0]})")
