import requests
import json
import openpyxl

BASE_URL = "http://localhost:5000/api"

print("=" * 70)
print("TESTING ARREAR WORKFLOW VIA OFFICIAL API ENDPOINTS")
print("=" * 70)

# 1. Fetch students to pick a test student
res = requests.get(f"{BASE_URL}/students")
students = res.json().get('students', [])
assert len(students) > 0, "No students found"

# Let's use test student (kalai)
test_std = [s for s in students if s['register_number'] == '113024104063']
student = test_std[0] if test_std else students[0]

std_id = student['id']
reg_num = student['register_number']
print(f"Selected Test Student: {student['name']} ({reg_num}) [ID: {std_id}]")

# STEP 1: Assign 'RA' to HS8151 in Semester 1 via /api/academic/save-grades
print("\n--- STEP 1: Save Grade 'RA' for HS8151 in Semester 1 ---")
save_payload = {
    "student_id": std_id,
    "semester_number": 1,
    "exam_month": "January",
    "exam_year": "2021",
    "regular_grades": [
        {
            "subject_code": "HS8151",
            "subject_name": "COMMUNICATIVE ENGLISH",
            "credits": 4,
            "grade": "RA"
        }
    ]
}
res_save = requests.post(f"{BASE_URL}/academic/save-grades", json=save_payload)
print(f"Save-grades response: {res_save.status_code} - {res_save.json()}")
assert res_save.status_code == 200

# STEP 2: Verify in /api/arrears (Website -> Pending Arrear)
print("\n--- STEP 2: Verify Pending Arrear in Website API ---")
res_arr = requests.get(f"{BASE_URL}/arrears?status=PENDING&student_id={reg_num}")
pending_arrears = res_arr.json().get('arrears', [])
print(f"Pending arrears count for student: {len(pending_arrears)}")
matching_arr = [a for a in pending_arrears if a['subject_code'] == 'HS8151']
assert len(matching_arr) > 0, "Pending arrear for HS8151 not found!"
arrear_record = matching_arr[0]
print(f"  [OK] Pending Arrear found: ID={arrear_record['id']}, Subject={arrear_record['subject_code']}, Grade={arrear_record['original_grade']}, Status={arrear_record['status']}")

# STEP 3: Verify Excel Export reflects RA
print("\n--- STEP 3: Verify Excel Export reflects Pending Arrear 'RA' ---")
res_excel = requests.get(f"{BASE_URL}/excel/export")
with open("test_export_pending.xlsx", "wb") as f:
    f.write(res_excel.content)

wb = openpyxl.load_workbook("test_export_pending.xlsx", data_only=True)
ws = wb['20-24']
std_row = None
for r in range(7, ws.max_row + 1):
    if str(ws.cell(r, 3).value).strip() == reg_num:
        std_row = r
        break

assert std_row is not None, f"Student {reg_num} not found in Excel!"
grade_val = ws.cell(std_row, 5).value  # Col 5 is HS8151
print(f"Excel Row {std_row}, Col 5 (HS8151): '{grade_val}'")
assert grade_val == 'RA', f"Expected 'RA', got '{grade_val}'"
print("  [OK] Excel export reflects pending 'RA' grade!")

# STEP 4: Clear the Arrear
print("\n--- STEP 4: Clear the Arrear via /api/arrears/clear ---")
clear_payload = {
    "arrear_id": arrear_record['id'],
    "cleared_semester": 2,
    "cleared_grade": "A",
    "cleared_month": "June",
    "cleared_year": "2021"
}
res_clear = requests.post(f"{BASE_URL}/arrears/clear", json=clear_payload)
print(f"Clear arrear response: {res_clear.status_code} - {res_clear.json()}")
assert res_clear.status_code == 200

# STEP 5: Verify Cleared Arrear & History in Website API
print("\n--- STEP 5: Verify Cleared Arrear & Arrear History ---")
# Cleared list
res_cleared = requests.get(f"{BASE_URL}/arrears?status=CLEARED&student_id={reg_num}")
cleared_list = res_cleared.json().get('arrears', [])
cleared_arr = [a for a in cleared_list if a['subject_code'] == 'HS8151'][0]
print(f"Cleared Record: ID={cleared_arr['id']}, OrigGrade={cleared_arr['original_grade']}, ClearedGrade={cleared_arr['cleared_grade']}, Sem={cleared_arr['cleared_semester']}, Status={cleared_arr['status']}")
assert cleared_arr['original_grade'] == 'RA', "Original RA grade history must not disappear!"
assert cleared_arr['cleared_grade'] == 'A', "Cleared grade must be A!"

# Student history timeline
res_hist = requests.get(f"{BASE_URL}/academic/student-history/{reg_num}")
hist_data = res_hist.json()
print(f"Student history arrears timeline records: {len(hist_data.get('arrears_timeline', []))}")
assert any(a['subject_code'] == 'HS8151' and a['status'] == 'CLEARED' for a in hist_data['arrears_timeline'])
print("  [OK] Original RA history preserved in Website Arrear History!")

# STEP 6: Verify in Excel Export
print("\n--- STEP 6: Verify Cleared Arrear in Excel Export ---")
res_excel_cleared = requests.get(f"{BASE_URL}/excel/export")
with open("test_export_cleared.xlsx", "wb") as f:
    f.write(res_excel_cleared.content)

wb_c = openpyxl.load_workbook("test_export_cleared.xlsx", data_only=False)
ws_c = wb_c['20-24']
cleared_cell = ws_c.cell(std_row, 5)
cleared_grade = cleared_cell.value
cleared_color = cleared_cell.fill.fgColor.rgb if cleared_cell.fill and cleared_cell.fill.fgColor else None
print(f"Excel Row {std_row}, Col 5: Grade='{cleared_grade}', Color='{cleared_color}'")
assert cleared_grade == 'A', f"Expected 'A', got '{cleared_grade}'"
# Semester 2 clearance color should be applied: FFB3CEFA
print(f"Expected Sem 2 Clearance Color: FFB3CEFA, Actual Color: {cleared_color}")
assert cleared_color == 'FFB3CEFA', f"Clearance color mismatch: {cleared_color}"
print("  [OK] Cleared grade 'A' and clearance semester color FFB3CEFA present in Excel!")

print("\n" + "=" * 70)
print("ALL ARREAR WORKFLOW TESTS PASSED PERFECTLY!")
print("=" * 70)
