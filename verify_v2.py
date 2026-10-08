import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ORIG   = r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx'
EXPORT = r'e:\COLLEGE PROJECT\test_export_v2.xlsx'

wb_orig   = openpyxl.load_workbook(ORIG, data_only=True)
wb_orig_f = openpyxl.load_workbook(ORIG, data_only=False)
wb_exp    = openpyxl.load_workbook(EXPORT, data_only=True)
wb_exp_f  = openpyxl.load_workbook(EXPORT, data_only=False)

ws_o = wb_orig['20-24']
ws_e = wb_exp['20-24']
ws_o_f = wb_orig_f['20-24']
ws_e_f = wb_exp_f['20-24']

print('=== DIMENSIONS ===')
print(f'Original: rows={ws_o.max_row}, cols={ws_o.max_column}')
print(f'Export:   rows={ws_e.max_row}, cols={ws_e.max_column}')

print('\n=== MERGED CELLS ===')
orig_merges = set(str(m) for m in ws_o_f.merged_cells.ranges)
exp_merges  = set(str(m) for m in ws_e_f.merged_cells.ranges)
missing = orig_merges - exp_merges
extra   = exp_merges  - orig_merges
print(f'Original: {len(orig_merges)}, Export: {len(exp_merges)}, Missing: {len(missing)}, Extra: {len(extra)}')
if missing:
    print('Missing (first 10):')
    for m in sorted(missing)[:10]: print(f'  {m}')

print('\n=== COLUMN POSITIONS (row 4) ===')
seen_o = {}
for c in range(1, ws_o.max_column + 1):
    v = ws_o.cell(4, c).value
    if v and str(v).strip():
        code = str(v).strip()
        if code not in seen_o:
            seen_o[code] = c

seen_e = {}
for c in range(1, ws_e.max_column + 1):
    v = ws_e.cell(4, c).value
    if v and str(v).strip():
        code = str(v).strip()
        if code not in seen_e:
            seen_e[code] = c

mismatches = []
for code, col in seen_o.items():
    if code in seen_e:
        if seen_e[code] != col:
            mismatches.append((code, col, seen_e[code]))
    else:
        mismatches.append((code, col, 'MISSING'))

print(f'Col mismatches: {len(mismatches)}')
if mismatches:
    for code, oc, ec in mismatches[:20]:
        print(f'  {code:<22} orig=col{oc:<6} export=col{ec}')
else:
    print('[OK] ALL columns in same positions')

print('\n=== GRADE VERIFICATION (Sem 1, first 5 students) ===')
sem1_grade_cols = [5, 7, 9, 11, 13, 15, 17, 19]
subj = ['HS8151','MA8151','PH8151','CY8151','GE8151','GE8152','GE8161','BS8161']
any_diff = False
for row in range(7, 12):
    reg = ws_o.cell(row, 3).value
    for i, col in enumerate(sem1_grade_cols):
        og = str(ws_o.cell(row, col).value or '').strip()
        eg = str(ws_e.cell(row, col).value or '').strip()
        ok = og == eg
        if not ok:
            any_diff = True
            print(f'  DIFF: R{row} {subj[i]:<10} orig={og!r:<8} export={eg!r:<8}')

if not any_diff:
    print('[OK] All Sem 1 grades match exactly (letter grades preserved)')

print('\n=== FORMULA CELLS (GP cols) ===')
for row in [7, 8]:
    for col in [6, 8, 10]:  # GP formula cols
        ov = ws_o_f.cell(row, col).value
        ev = ws_e_f.cell(row, col).value
        ok = type(ov) == type(ev) or (isinstance(ov, str) and ov.startswith('=') and isinstance(ev, str) and ev.startswith('='))
        print(f'  R{row}C{col}: orig={repr(str(ov)[:40])}, export={repr(str(ev)[:40])} {"[OK]" if ok else "[DIFF]"}')

print('\n=== HEADER ROWS 1-6 ===')
hm = 0
for r in range(1, 7):
    for c in range(1, min(ws_o.max_column + 1, 180)):
        o = str(ws_o.cell(r, c).value or '').strip()
        e = str(ws_e.cell(r, c).value or '').strip()
        if o != e:
            hm += 1
            if hm <= 5:
                print(f'  DIFF R{r}C{c}: orig={repr(o)[:25]}, export={repr(e)[:25]}')
print(f'Total header mismatches: {hm}' if hm else '[OK] All headers match')

print('\n=== STUDENT COUNT ===')
orig_studs = sum(1 for r in range(7, ws_o.max_row+1)
                 if ws_o.cell(r,3).value and str(ws_o.cell(r,3).value).strip().isdigit() and len(str(ws_o.cell(r,3).value).strip()) >= 8)
exp_studs  = sum(1 for r in range(7, ws_e.max_row+1)
                 if ws_e.cell(r,3).value and str(ws_e.cell(r,3).value).strip().isdigit() and len(str(ws_e.cell(r,3).value).strip()) >= 8)
print(f'Original Excel students: {orig_studs}')
print(f'Exported Excel students: {exp_studs}')
