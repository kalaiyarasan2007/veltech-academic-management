import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)

for sname in ['I API', 'II API ', 'III API']:
    ws = wb[sname]
    print(f"\n=================== SHEET: '{sname}' ===================")
    for r in range(1, 20):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {[v for v in row_vals if v is not None]}")
