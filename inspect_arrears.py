import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']

# Let's check grade columns across all semesters for each student
# Grade columns are:
# Sem 1: E, G, I, K, M, O, Q, S (Cols 5, 7, 9, 11, 13, 15, 17, 19)
# Sem 2: Y, AA, AC, AE, AG, AI, AK, AM (Cols 25, 27, 29, 31, 33, 35, 37, 39)
# Sem 3: AU, AW, AY, BA, BC, BE, BG, BI, BK (Cols 47, 49, 51, 53, 55, 57, 59, 61, 63)
# Sem 4: BS, BU, BW, BY, CA, CC, CE, CG, CI (Cols 71, 73, 75, 77, 79, 81, 83, 85, 87)
# Sem 5: CQ, CS, CU, CW, CY, DA, DC, DE, DG (Cols 95, 97, 99, 101, 103, 105, 107, 109, 111)
# Sem 6: DO, DQ, DS, DU, DW, DY, EA, EC, EE, EG (Cols 119, 121, 123, 125, 127, 129, 131, 133, 135, 137)
# Sem 7: EO, EQ, ES, EU, EW, EY, FA, FC (Cols 145, 147, 149, 151, 153, 155, 157, 159)
# Sem 8: FK, FM, FO (Cols 167, 169, 171)

sem_grade_cols = {
    1: [5, 7, 9, 11, 13, 15, 17, 19],
    2: [25, 27, 29, 31, 33, 35, 37, 39],
    3: [47, 49, 51, 53, 55, 57, 59, 61, 63],
    4: [71, 73, 75, 77, 79, 81, 83, 85, 87],
    5: [95, 97, 99, 101, 103, 105, 107, 109, 111],
    6: [119, 121, 123, 125, 127, 129, 131, 133, 135, 137],
    7: [145, 147, 149, 151, 153, 155, 157, 159],
    8: [167, 169, 171]
}

sem_summary_cols = {
    1: {'total': 21, 'gpa': 22, 'arr_sem': 23, 'arr_cum': None},
    2: {'total': 41, 'gpa': 42, 'cgpa': 43, 'arr_sem': 44, 'arr_cum': 45},
    3: {'total': 65, 'gpa': 66, 'cgpa': 67, 'arr_sem': 68, 'arr_cum': 69},
    4: {'total': 89, 'gpa': 90, 'cgpa': 91, 'arr_sem': 92, 'arr_cum': 93},
    5: {'total': 113, 'gpa': 114, 'cgpa': 115, 'arr_sem': 116, 'arr_cum': 117},
    6: {'total': 139, 'gpa': 140, 'cgpa': 141, 'arr_sem': 142, 'arr_cum': 143},
    7: {'total': 161, 'gpa': 162, 'cgpa': 163, 'arr_sem': 164, 'arr_cum': 165},
    8: {'total': 173, 'gpa': 174, 'cgpa': 175, 'arr_sem': 176, 'arr_cum': 177}
}

students_with_arrears = []

for r in range(7, 133):
    reg = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value else ""
    name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value else ""
    if not reg or not reg.isdigit():
        continue
    
    student_arrears = []
    for sem, cols in sem_grade_cols.items():
        for c in cols:
            val = str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else ""
            subj_code = str(ws.cell(4, c).value).strip()
            # check cell fill color or grade value ('U', 'RA', 'AB', 'WH', 'WD', etc.)
            cell_fill = ws.cell(r, c).fill
            fg = cell_fill.fgColor if cell_fill else None
            rgb = fg.rgb if (fg and fg.type == 'rgb') else None
            
            if val in ['RA', 'U', 'AB', 'WH', 'WD'] or rgb is not None:
                student_arrears.append((sem, c, subj_code, val, rgb))
    
    if student_arrears:
        students_with_arrears.append((r, reg, name, student_arrears))

print(f"Total students with arrears/special fills: {len(students_with_arrears)}")
for r, reg, name, arrs in students_with_arrears[:10]:
    print(f"\nStudent Row {r}: {reg} - {name}")
    for sem, c, code, val, rgb in arrs:
        print(f"  Sem {sem} (Col {c} {openpyxl.utils.get_column_letter(c)}): Code={code}, Grade={val}, Color={rgb}")
