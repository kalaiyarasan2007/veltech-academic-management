import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=False)
ws = wb['20-24']

print("=== MERGED RANGES IN 20-24 ===")
for m in list(ws.merged_cells.ranges)[:30]:
    print(m)

print("\n=== HEADER ROWS 1 TO 6 FOR COLUMNS 1 TO 178 ===")
for col in range(1, ws.max_column + 1):
    c_let = openpyxl.utils.get_column_letter(col)
    cell_r4 = ws.cell(4, col)
    cell_r5 = ws.cell(5, col)
    cell_r6 = ws.cell(6, col)
    cell_r7 = ws.cell(7, col)
    
    # check fill on r1..r6
    fills = []
    for r in range(1, 7):
        cell = ws.cell(r, col)
        if cell.fill and cell.fill.fill_type:
            fg = cell.fill.fgColor
            if fg:
                if fg.type == 'rgb':
                    fills.append(f"R{r}:{fg.rgb}")
                elif fg.type == 'theme':
                    fills.append(f"R{r}:theme{fg.theme}_tint{fg.tint}")
                elif fg.indexed is not None:
                    fills.append(f"R{r}:idx{fg.indexed}")
    
    val4 = str(cell_r4.value).replace('\n', ' ') if cell_r4.value else ""
    val5 = str(cell_r5.value).replace('\n', ' ') if cell_r5.value else ""
    val6 = str(cell_r6.value).replace('\n', ' ') if cell_r6.value else ""
    val7 = str(cell_r7.value).replace('\n', ' ') if cell_r7.value else ""
    
    if val4 or val5 or val6:
        print(f"Col {col:3d} ({c_let:3s}) | Fills: {', '.join(fills):30s} | R4: {val4[:25]:25s} | R5: {val5[:25]:25s} | R6: {val6:4s} | R7: {val7[:35]}")
