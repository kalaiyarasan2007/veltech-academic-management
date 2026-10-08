import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']

# Find all columns that have subject code in row 4, title in row 5, credit in row 6
subjects_by_sem = {1: [], 2: [], 3: [], 4: [], 5: [], 6: [], 7: [], 8: []}

# Let's map column ranges to semesters from row 4 headers and summary columns
# Total grade / GPA / Arrear columns mark the end of each semester!
# Let's scan all columns 1..178

current_sem = 1
for c in range(1, ws.max_column + 1):
    c_let = openpyxl.utils.get_column_letter(c)
    val4 = ws.cell(4, c).value
    val5 = ws.cell(5, c).value
    val6 = ws.cell(6, c).value
    
    val4_str = str(val4).strip() if val4 else ""
    val5_str = str(val5).strip() if val5 else ""
    val6_str = str(val6).strip() if val6 else ""
    
    if "Arrear" in val4_str or "Arrear" in val5_str:
        if "1st" in val4_str or "1st" in val5_str or "Ist" in val4_str:
            current_sem = 2
        elif "2nd" in val4_str or "Ist year" in val4_str:
            current_sem = 3
        elif "3rd" in val4_str:
            current_sem = 4
        elif "4th" in val4_str:
            current_sem = 5
        elif "5th" in val4_str:
            current_sem = 6
        elif "6th" in val4_str:
            current_sem = 7
        elif "7th" in val4_str:
            current_sem = 8
    
    # Check if this column is a subject (has subject code in R4, name in R5, credit in R6)
    if val4_str and val5_str and val6_str.isdigit():
        credit = int(val6_str)
        subjects_by_sem[current_sem].append({
            'col_num': c,
            'col_let': c_let,
            'code': val4_str,
            'title': val5_str,
            'credit': credit
        })

print("=== EXTRACTED SUBJECTS FROM EXCEL '20-24' ===")
for sem, subs in subjects_by_sem.items():
    tot_credits = sum(s['credit'] for s in subs)
    print(f"\n--- SEMESTER {sem} (Total Credits = {tot_credits}) ---")
    for s in subs:
        print(f"  Col {s['col_num']:3d} ({s['col_let']:3s}): {s['code']} - {s['title']} - Credits: {s['credit']}")
