import openpyxl

# Load master Excel template
template_path = r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx'
output_path = r'e:\COLLEGE PROJECT\CN_PROJECT_UPDATED.xlsx'

print("Loading master Excel template:", template_path)
wb = openpyxl.load_workbook(template_path, data_only=False)

print("Master Sheet Names:", wb.sheetnames)
ws_main = wb['20-24']
print("Main sheet max_row:", ws_main.max_row, "max_column:", ws_main.max_column)

# Verify student lookup by Register Number in Col 3 (C)
student_rows = {}
for r in range(7, ws_main.max_row + 1):
    reg_val = ws_main.cell(r, 3).value
    if reg_val:
        reg_str = str(reg_val).strip()
        if reg_str.isdigit() and len(reg_str) >= 8:
            student_rows[reg_str] = r

print(f"Mapped {len(student_rows)} existing student rows in master template.")
print("Sample Reg 113020104001 is at Row:", student_rows.get('113020104001'))

# Save a copy to verify structure preservation
wb.save(output_path)
print("Saved template copy to:", output_path)
