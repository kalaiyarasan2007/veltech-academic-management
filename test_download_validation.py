import urllib.request
import openpyxl
import io

url = 'http://localhost:5000/api/excel/export'
output_filename = 'Academic_Records_Updated.xlsx'

print(f"Downloading from {url}...")
req = urllib.request.Request(url)
with urllib.request.urlopen(req) as response:
    content_type = response.headers.get('Content-Type')
    content_disp = response.headers.get('Content-Disposition')
    data = response.read()

    print(f"HTTP Status: {response.status}")
    print(f"Content-Type: {content_type}")
    print(f"Content-Disposition: {content_disp}")
    print(f"File Size: {len(data)} bytes")

    with open(output_filename, 'wb') as f:
        f.write(data)

print(f"Saved binary file to disk as '{output_filename}'.")

# Validate file using openpyxl
print("\n--- VALIDATING EXCEL WORKBOOK CONTENT ---")
wb = openpyxl.load_workbook(output_filename, data_only=True)
print("Worksheet Names:", wb.sheetnames)

ws_main = wb['20-24']
print("Main Sheet '20-24' Max Rows:", ws_main.max_row, "Max Cols:", ws_main.max_column)

student_reg = ws_main.cell(7, 3).value
student_name = ws_main.cell(7, 4).value

print(f"Student Row 7: Reg No = {student_reg}, Name = {student_name}")

assert len(wb.sheetnames) == 4, "Expected 4 worksheets"
assert '20-24' in wb.sheetnames, "Missing '20-24' sheet"
assert str(student_reg) == '113020104001', "Student Reg No mismatch"

print("\nSUCCESS: Valid Academic_Records_Updated.xlsx file created and verified!")
