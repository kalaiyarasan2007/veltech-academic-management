import urllib.request
import urllib.parse
import json
import openpyxl
import subprocess
import os
import sys

BASE_URL = "http://localhost:5000/api"

print("=================================================================")
print("STARTING 30-POINT FINAL ACCEPTANCE TEST")
print("=================================================================")

def api_get(endpoint):
    url = f"{BASE_URL}{endpoint}"
    req = urllib.request.Request(url)
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode('utf-8'))

def api_post(endpoint, data):
    url = f"{BASE_URL}{endpoint}"
    req = urllib.request.Request(
        url,
        data=json.dumps(data).encode('utf-8'),
        headers={'Content-Type': 'application/json'}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode('utf-8'))

def api_download(endpoint, out_path):
    url = f"{BASE_URL}{endpoint}"
    req = urllib.request.Request(url)
    with urllib.request.urlopen(req) as resp:
        headers = {k.lower(): v for k, v in resp.headers.items()}
        content = resp.read()
        with open(out_path, 'wb') as f:
            f.write(content)
        return headers, len(content)

# TEST 1 & 2: Add / Update student
test_reg = "113024104099"
test_name = "ANANDHA KRISHNAN M"
test_roll = "20CS999"

print("\n--- TEST 1 & 2: Add Student ---")
add_resp = api_post("/students", {
    "register_number": test_reg,
    "name": test_name,
    "roll_number": test_roll,
    "batch": "2020-2024",
    "department": "Department of Computer Science and Engineering"
})
print("Added/Updated student:", add_resp.get("student", {}).get("name"), add_resp.get("student", {}).get("register_number"))
student_id = add_resp.get("student", {}).get("id")
assert student_id, "Failed to get student ID"

# TEST 3, 4, 5: Semester 1 grades with exam month/year
print("\n--- TEST 3, 4, 5: Save Semester 1 Grades ---")
sem1_grades = [
    {"subject_code": "HS8151", "subject_name": "Communicative English", "credits": 4, "grade": "A"},
    {"subject_code": "MA8151", "subject_name": "Engineering Maths I", "credits": 4, "grade": "O"},
    {"subject_code": "PH8151", "subject_name": "Engineering Physics", "credits": 3, "grade": "A+"},
    {"subject_code": "CY8151", "subject_name": "Engineering Chemistry", "credits": 3, "grade": "A"},
    {"subject_code": "GE8151", "subject_name": "Python Pgmming", "credits": 3, "grade": "A"},
    {"subject_code": "GE8152", "subject_name": "Engineering Graphics", "credits": 4, "grade": "B+"},
    {"subject_code": "GE8161", "subject_name": "Python Pgmming Laboratory", "credits": 2, "grade": "O"},
    {"subject_code": "BS8161", "subject_name": "Physics and Chemistry Laboratory", "credits": 2, "grade": "O"}
]
# Expected GP:
# HS8151: A (8) * 4 = 32
# MA8151: O (10) * 4 = 40
# PH8151: A+ (9) * 3 = 27
# CY8151: A (8) * 3 = 24
# GE8151: A (8) * 3 = 24
# GE8152: B+ (7) * 4 = 28
# GE8161: O (10) * 2 = 20
# BS8161: O (10) * 2 = 20
# Total grade = 32+40+27+24+24+28+20+20 = 215
# Credits = 25
# GPA = 215 / 25 = 8.60

save1_resp = api_post("/academic/save-grades", {
    "student_id": test_reg,
    "semester_number": 1,
    "exam_month": "January",
    "exam_year": "2021",
    "regular_grades": sem1_grades
})
print("Saved Sem 1 grades response:", save1_resp.get("message", "OK"))

# TEST 6, 7, 8: Verify Grade Points, GPA, CGPA
print("\n--- TEST 6, 7, 8: Verify Sem 1 Academic Summary ---")
history1 = api_get(f"/academic/student-history/{test_reg}")
sem1_summary = history1["semesters"][0]
print(f"Sem 1 Total Grade: {sem1_summary['total_grade']} (Expected: 215)")
print(f"Sem 1 GPA: {sem1_summary['gpa']} (Expected: 8.60)")
print(f"Sem 1 CGPA: {sem1_summary['cgpa']} (Expected: 8.60)")
assert sem1_summary['total_grade'] == 215, f"Expected 215, got {sem1_summary['total_grade']}"
assert sem1_summary['gpa'] == 8.60, f"Expected 8.60, got {sem1_summary['gpa']}"

# TEST 9 & 10: Create an RA in Semester 2 and Verify Pending Arrear
print("\n--- TEST 9 & 10: Create RA in Semester 2 ---")
sem2_grades = [
    {"subject_code": "HS8251", "subject_name": "Technical English", "credits": 4, "grade": "A"},
    {"subject_code": "MA8251", "subject_name": "Engineering Mathmatics II", "credits": 4, "grade": "RA"}, # RA arrear!
    {"subject_code": "PH 8252", "subject_name": "Physics and Information Science", "credits": 3, "grade": "A"},
    {"subject_code": "BE8255", "subject_name": "BEEE", "credits": 3, "grade": "B+"},
    {"subject_code": "GE8291", "subject_name": "EVS", "credits": 3, "grade": "A"},
    {"subject_code": "CS8251", "subject_name": "Programming in c", "credits": 3, "grade": "A"},
    {"subject_code": "GE8261", "subject_name": "Engineering Practices Lab", "credits": 2, "grade": "O"},
    {"subject_code": "CS8261", "subject_name": "C Pgmming Lab", "credits": 2, "grade": "O"}
]
save2_resp = api_post("/academic/save-grades", {
    "student_id": test_reg,
    "semester_number": 2,
    "exam_month": "June",
    "exam_year": "2021",
    "regular_grades": sem2_grades
})
print("Saved Sem 2 grades with RA.")

arrears_resp = api_get(f"/arrears?student_id={test_reg}")
arrears_list = arrears_resp.get("arrears", [])
print(f"Found {len(arrears_list)} arrears for {test_reg}:")
for a in arrears_list:
    print(f"  Arrear: {a['subject_code']} - {a['subject_name']}, Status: {a['status']}, Orig Grade: {a['original_grade']}")
assert len(arrears_list) >= 1, "Arrear was not created!"
ma_arrear = [a for a in arrears_list if a['subject_code'] == 'MA8251'][0]
assert ma_arrear['status'] == 'PENDING', "Arrear status should be PENDING"

# TEST 11, 12, 13, 14, 15: Clear Arrear in Semester 3
print("\n--- TEST 11, 12, 13, 14, 15: Clear Arrear ---")
clear_resp = api_post("/arrears/clear", {
    "arrear_id": ma_arrear['id'],
    "cleared_semester": 3,
    "cleared_grade": "A",
    "cleared_month": "January",
    "cleared_year": "2022"
})
print("Cleared arrear response:", clear_resp.get("message"))

arrears_resp_after = api_get(f"/arrears?student_id={test_reg}")
ma_arrear_after = [a for a in arrears_resp_after.get("arrears", []) if a['subject_code'] == 'MA8251'][0]
print(f"Arrear after clearance: Status={ma_arrear_after['status']}, Cleared Grade={ma_arrear_after['cleared_grade']}, Cleared Sem={ma_arrear_after['cleared_semester']}")
assert ma_arrear_after['status'] == 'CLEARED', "Arrear should now be CLEARED"
assert ma_arrear_after['cleared_grade'] == 'A', "Cleared grade should be A"

# TEST 16 & 17: Generate Updated Excel & Preview
print("\n--- TEST 16 & 17: Excel Preview ---")
preview_data = api_get("/excel/preview")
sheets = preview_data.get("sheets", [])
print(f"Preview returned {len(sheets)} sheets: {[s['name'] for s in sheets]}")
assert len(sheets) == 4, f"Expected 4 sheets, got {len(sheets)}"

# TEST 18, 19, 20, 21, 22: Verify Preview Details
print("\n--- TEST 18, 19, 20, 21, 22: Verify Preview Content ---")
main_preview = [s for s in sheets if s['name'] == '20-24'][0]
row1_cells = [c['value'] for c in main_preview['rows'][0]['cells'] if c['value']]
print("Preview Row 1 Heading:", row1_cells[0][:60] if row1_cells else "None")
assert any("Vel Tech" in v for v in row1_cells), "Original college heading missing in preview!"

# Find test student in preview
test_student_row = None
for r in main_preview['rows']:
    c3 = [c['value'] for c in r['cells'] if c['col'] == 3]
    if c3 and c3[0] == test_reg:
        test_student_row = r
        break
assert test_student_row, f"Student {test_reg} not found in Excel Preview!"
print(f"Found student {test_reg} at Row {test_student_row['rowNumber']} in preview.")

# Check Sem 1 GPA and Sem 2 Cleared Arrear in preview
sem1_gpa_val = [c['value'] for c in test_student_row['cells'] if c['col'] == 22][0]
sem2_ma8251_val = [c['value'] for c in test_student_row['cells'] if c['col'] == 27][0]
sem2_ma8251_colr = [c['color'] for c in test_student_row['cells'] if c['col'] == 27][0]
print(f"Preview Sem 1 GPA (Col 22): {sem1_gpa_val}")
print(f"Preview MA8251 Grade (Col 27): {sem2_ma8251_val} with clearance color {sem2_ma8251_colr}")
assert sem1_gpa_val != "", "GPA value is empty in preview!"
assert sem2_ma8251_val == "A", f"Expected cleared grade A, got {sem2_ma8251_val}"
# Sem 3 color is #FFA766
assert sem2_ma8251_colr == "#FFA766", f"Expected clearance sem 3 color #FFA766, got {sem2_ma8251_colr}"

# TEST 23 & 24: Download Excel File
print("\n--- TEST 23 & 24: Download Updated Excel ---")
download_path = "CN_PROJECT_UPDATED_VERIFY.xlsx"
headers, size = api_download("/excel/export", download_path)
print(f"Downloaded file: size={size} bytes")
print("Content-Type header:", headers.get("content-type"))
print("Content-Disposition header:", headers.get("content-disposition"))
assert "spreadsheetml.sheet" in headers.get("content-type", ""), "Wrong content type"
assert "CN_PROJECT_UPDATED.xlsx" in headers.get("content-disposition", ""), "Wrong filename header"
assert size > 200000, f"File size too small: {size}"

# TEST 25 & 26: Open Downloaded File in Real Microsoft Excel via COM
print("\n--- TEST 25 & 26: Open Downloaded XLSX in Real Microsoft Excel ---")
ps_script = f"""
try {{
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $filePath = (Resolve-Path ".\\{download_path}").Path
    $wb = $excel.Workbooks.Open($filePath)
    $sheetCount = $wb.Sheets.Count
    $ws = $wb.Sheets.Item("20-24")
    $reg = $ws.Range("C{test_student_row['rowNumber']}").Text
    $name = $ws.Range("D{test_student_row['rowNumber']}").Text
    $sem1_gpa = $ws.Range("V{test_student_row['rowNumber']}").Text
    $sem2_gpa = $ws.Range("AP{test_student_row['rowNumber']}").Text
    $sem2_cgpa = $ws.Range("AQ{test_student_row['rowNumber']}").Text
    $ma8251_val = $ws.Range("AA{test_student_row['rowNumber']}").Text
    $wb.Close($false)
    $excel.Quit()
    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null
    Write-Host "EXCEL_COM_SUCCESS: Sheets=$sheetCount Reg=$reg Name=$name Sem1GPA=$sem1_gpa Sem2GPA=$sem2_gpa Sem2CGPA=$sem2_cgpa MA8251=$ma8251_val"
}} catch {{
    Write-Host "EXCEL_COM_ERROR: $($_.Exception.Message)"
}}
"""
with open("verify_com.ps1", "w", encoding="utf-8") as f:
    f.write(ps_script)

res = subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", "verify_com.ps1"], capture_output=True, text=True)
print(res.stdout)
assert "EXCEL_COM_SUCCESS" in res.stdout, f"Microsoft Excel failed to open workbook: {res.stdout} {res.stderr}"

# TEST 27, 28, 29, 30: Verify Structure, Sheets, Students, Website Details
print("\n--- TEST 27, 28, 29, 30: Detailed Verification ---")
wb_v = openpyxl.load_workbook(download_path, data_only=True)
assert wb_v.sheetnames == ['20-24', 'I API', 'II API ', 'III API'], f"Sheet names mismatch: {wb_v.sheetnames}"
ws_v = wb_v['20-24']
print(f"Worksheet '20-24' dimensions: {ws_v.max_row} rows x {ws_v.max_column} cols")
assert any(ws_v.cell(1, c).value and "Vel Tech" in str(ws_v.cell(1, c).value) for c in range(1, 20)), "Heading in row 1 missing"
assert ws_v.cell(4, 2).value == "S.No", "Col 2 in row 4 should be S.No"
assert "Register" in str(ws_v.cell(4, 3).value), "Col 3 in row 4 should be Register Number"

# Check our test student
r_num = test_student_row['rowNumber']
assert str(ws_v.cell(r_num, 3).value).strip() == test_reg, "Student Register number mismatch"
assert str(ws_v.cell(r_num, 4).value).strip() == test_name, "Student Name mismatch"
print(f"Verified Student in Excel: {ws_v.cell(r_num, 3).value} - {ws_v.cell(r_num, 4).value}")

# Check GPA and CGPA
print(f"Col 22 (Sem 1 GPA): {ws_v.cell(r_num, 22).value}")
print(f"Col 42 (Sem 2 GPA): {ws_v.cell(r_num, 42).value}")
print(f"Col 43 (Sem 2 CGPA): {ws_v.cell(r_num, 43).value}")
assert float(ws_v.cell(r_num, 22).value) == 8.6, f"Expected 8.6, got {ws_v.cell(r_num, 22).value}"

# Check cleared arrear grade
print(f"Col 27 (MA8251 cleared grade): {ws_v.cell(r_num, 27).value}")
assert ws_v.cell(r_num, 27).value == "A", f"Expected A, got {ws_v.cell(r_num, 27).value}"

# Check API Sheets
ws_api1 = wb_v['I API']
api_found = False
for r in range(10, ws_api1.max_row + 1):
    if str(ws_api1.cell(r, 3).value).strip() == test_reg:
        api_found = True
        print(f"Found student in 'I API' at Row {r}: SNo={ws_api1.cell(r, 2).value}, Sem1Tot={ws_api1.cell(r, 5).value}, Sem2Tot={ws_api1.cell(r, 7).value}, YearGPA={ws_api1.cell(r, 9).value}")
        break
assert api_found, "Student not found in I API sheet!"

print("\n=================================================================")
print("ALL 30 TEST POINTS PASSED 100%! VERIFICATION COMPLETE & SUCCESSFUL!")
print("=================================================================")
