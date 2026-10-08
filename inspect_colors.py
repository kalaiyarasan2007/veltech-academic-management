import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=False)
ws = wb['20-24']

print("=== INSPECTING COLORS IN 20-24 ===")
for r in range(1, 10):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(r, c)
        if cell.fill and cell.fill.fill_type:
            fg = cell.fill.fgColor
            bg = cell.fill.bgColor
            val = str(cell.value).replace('\n', ' ') if cell.value else ""
            print(f"R{r}C{c} ({openpyxl.utils.get_column_letter(c)}{r}): fill_type={cell.fill.fill_type}, fg_type={fg.type if fg else None}, fg_rgb={fg.rgb if fg else None}, fg_theme={fg.theme if fg else None}, fg_tint={fg.tint if fg else None}, val={val[:30]}")
