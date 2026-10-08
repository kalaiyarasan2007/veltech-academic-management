import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']

sem_colors = {
    1: 'FF92D050', # Green
    2: 'FFB3CEFA', # Soft Blue
    3: 'FFFFA766', # Soft Orange
    4: 'FF00B0F0', # Sky Blue
    5: 'FFFFFF00', # Yellow
    6: 'FFF28E85', # Soft Coral
    7: 'FFCC9900', # Gold
    8: 'NONE'
}
color_to_sem = {v: k for k, v in sem_colors.items() if v != 'NONE'}

sem_grade_cols = {
    1: [5, 7, 9, 11, 13, 15, 17, 19],
    2: [25, 27, 29, 31, 33, 35, 37, 39],
    3: [47, 49, 51, 53, 55, 57, 59, 61, 63],
    4: [71, 73, 75, 77, 79, 81, 83, 85, 87],
    5: [95, 97, 99, 101, 103, 105, 107, 109, 111],
    6: [119, 121, 123, 125, 127, 129, 131, 133, 135, 137],
    7: [145, 147, 149, 151, 153, 155, 157, 159],
    8: [167, 169, 171]
}

print("=== CLEAR ARREARS COLOR MISMATCHES (ORIGINAL SEM vs CLEARED SEM COLOR) ===")
count = 0
for r in range(7, 130):
    reg = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value else ""
    name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value else ""
    if not reg or not reg.isdigit() or len(reg) < 8:
        continue
    
    for orig_sem, cols in sem_grade_cols.items():
        expected_color = sem_colors[orig_sem]
        for c in cols:
            val = str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else ""
            subj_code = str(ws.cell(4, c).value).strip()
            
            cell_fill = ws.cell(r, c).fill
            fg = cell_fill.fgColor if cell_fill else None
            actual_color = fg.rgb if (fg and fg.type == 'rgb') else 'NONE'
            
            if actual_color in color_to_sem and color_to_sem[actual_color] != orig_sem:
                cleared_sem = color_to_sem[actual_color]
                print(f"Row {r:3d} | {reg} | {name[:20]:20s} | Subj: {subj_code:8s} | Orig Sem {orig_sem} | Cleared Grade: {val:4s} | Cleared in Sem {cleared_sem} (Color: {actual_color})")
                count += 1

print(f"\nTotal cleared arrear color mismatches found: {count}")
