import openpyxl

wb = openpyxl.load_workbook(r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx', data_only=True)
ws = wb['20-24']

print("=== CHECKING ROWS BELOW STUDENT LIST (Row 130 to 332) ===")
for r in range(130, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 25)]
    non_empty = [v for v in vals if v is not None and str(v).strip() != ""]
    if non_empty:
        print(f"Row {r:3d}: {non_empty[:10]}")
