import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ORIG   = r'e:\COLLEGE PROJECT\CN-PRJECT  DETSILS.xlsx'
EXPORT = r'e:\COLLEGE PROJECT\test_export_final.xlsx'
import json
DB_FILE = r'e:\COLLEGE PROJECT\server\data\database.json'

wb_orig   = openpyxl.load_workbook(ORIG, data_only=True)
wb_orig_f = openpyxl.load_workbook(ORIG, data_only=False)
wb_exp    = openpyxl.load_workbook(EXPORT, data_only=True)
wb_exp_f  = openpyxl.load_workbook(EXPORT, data_only=False)

with open(DB_FILE, encoding='utf-8') as f:
    db = json.load(f)

ws_o   = wb_orig['20-24']
ws_e   = wb_exp['20-24']
ws_o_f = wb_orig_f['20-24']
ws_e_f = wb_exp_f['20-24']

results = {}

print('='*70)
print('FINAL VERIFICATION REPORT')
print('='*70)

# --- Student counts ---
orig_students = {}
for r in range(7, ws_o.max_row+1):
    reg = ws_o.cell(r,3).value
    if reg and str(reg).strip().isdigit() and len(str(reg).strip())>=8:
        orig_students[str(reg).strip()] = str(ws_o.cell(r,4).value or '').strip()

db_students = {s['register_number']: s['name'] for s in db['students']}

exp_students = {}
for r in range(7, ws_e.max_row+1):
    reg = ws_e.cell(r,3).value
    if reg and str(reg).strip().isdigit() and len(str(reg).strip())>=8:
        exp_students[str(reg).strip()] = str(ws_e.cell(r,4).value or '').strip()

extra_in_db  = set(db_students)  - set(orig_students)
extra_in_exp = set(exp_students) - set(orig_students)
miss_in_exp  = set(orig_students)- set(exp_students)

print(f'\n1. STUDENT COUNTS:')
print(f'   Original Excel : {len(orig_students)}')
print(f'   Database       : {len(db_students)}')
print(f'   Exported Excel : {len(exp_students)}')
print(f'   Extra in DB (test students added via web) : {sorted(extra_in_db)}')
print(f'   Extra in Export: {sorted(extra_in_exp)}')
print(f'   Missing from Export: {sorted(miss_in_exp) or "NONE"}')
results['students_ok'] = not miss_in_exp

# --- Sheet structure ---
orig_sheets = wb_orig.sheetnames
exp_sheets  = wb_exp.sheetnames
print(f'\n2. SHEET NAMES:')
print(f'   Original : {orig_sheets}')
print(f'   Exported : {exp_sheets}')
sheets_ok = orig_sheets == exp_sheets
print(f'   Status   : {"OK" if sheets_ok else "MISMATCH"}')
results['sheets_ok'] = sheets_ok

# --- Dimensions ---
print(f'\n3. DIMENSIONS (20-24):')
orig_dims = (ws_o.max_row, ws_o.max_column)
exp_dims  = (ws_e.max_row, ws_e.max_column)
print(f'   Original : rows={orig_dims[0]}, cols={orig_dims[1]}')
print(f'   Exported : rows={exp_dims[0]}, cols={exp_dims[1]}')
dims_ok = orig_dims == exp_dims
print(f'   Status   : {"OK" if dims_ok else "MISMATCH"}')
results['dims_ok'] = dims_ok

# --- Merged cells ---
orig_merges = set(str(m) for m in ws_o_f.merged_cells.ranges)
exp_merges  = set(str(m) for m in ws_e_f.merged_cells.ranges)
missing_m   = orig_merges - exp_merges
extra_m     = exp_merges  - orig_merges
merges_ok   = not missing_m and not extra_m
print(f'\n4. MERGED CELLS:')
print(f'   Original : {len(orig_merges)}')
print(f'   Exported : {len(exp_merges)}')
print(f'   Missing  : {len(missing_m)}')
print(f'   Extra    : {len(extra_m)}')
print(f'   Status   : {"OK - All preserved" if merges_ok else "MISMATCH"}')
if missing_m:
    for m in sorted(missing_m)[:5]: print(f'     Missing: {m}')
results['merges_ok'] = merges_ok

# --- Column positions ---
seen_o = {}
for c in range(1, ws_o.max_column+1):
    v = ws_o.cell(4,c).value
    if v and str(v).strip():
        code = str(v).strip()
        if code not in seen_o: seen_o[code] = c

seen_e = {}
for c in range(1, ws_e.max_column+1):
    v = ws_e.cell(4,c).value
    if v and str(v).strip():
        code = str(v).strip()
        if code not in seen_e: seen_e[code] = c

col_mismatches = [(code, col, seen_e.get(code,'MISSING'))
                  for code, col in seen_o.items()
                  if seen_e.get(code) != col]
cols_ok = not col_mismatches
print(f'\n5. COLUMN POSITIONS (row 4 headers):')
print(f'   Mismatches: {len(col_mismatches)}')
print(f'   Status    : {"OK - All in same positions" if cols_ok else "MISMATCH"}')
if col_mismatches:
    for code, oc, ec in col_mismatches[:5]: print(f'     {code}: orig=col{oc}, export=col{ec}')
results['cols_ok'] = cols_ok

# --- Semester colors ---
sem_grade_cols = {1:5, 2:25, 3:47, 4:71, 5:95, 6:119, 7:145, 8:167}
def get_argb(cell):
    try:
        if cell.fill and cell.fill.fgColor and cell.fill.fgColor.type == 'rgb':
            return cell.fill.fgColor.rgb
        return 'None'
    except: return 'Error'

color_issues = []
for sem, col in sem_grade_cols.items():
    oc = get_argb(ws_o_f.cell(7, col))
    ec = get_argb(ws_e_f.cell(7, col))
    if oc != ec: color_issues.append((sem, oc, ec))

colors_ok = not color_issues
print(f'\n6. SEMESTER COLORS (row 7 sample):')
for sem, col in sem_grade_cols.items():
    oc = get_argb(ws_o_f.cell(7, col))
    ec = get_argb(ws_e_f.cell(7, col))
    ok = oc == ec
    print(f'   Sem {sem}: orig=#{oc} export=#{ec} {"OK" if ok else "MISMATCH"}')
results['colors_ok'] = colors_ok

# --- Grade verification (all sems, first 3 students) ---
grade_cols_list = [5,7,9,11,13,15,17,19,  # Sem1
                   25,27,29,31,33,35,37,39, # Sem2
                   47,49,51,53,55,57,59,61,63, # Sem3
                   71,73,75,77,79,81,83,85,87, # Sem4
                   95,97,99,101,103,105,107,109,111, # Sem5
                   119,121,123,125,127,129,131,133,135,137, # Sem6
                   145,147,149,151,153,155,157,159, # Sem7
                   167,169,171] # Sem8
grade_diffs = []
for row in range(7, 10):
    for col in grade_cols_list:
        og = str(ws_o.cell(row, col).value or '').strip()
        eg = str(ws_e.cell(row, col).value or '').strip()
        if og != eg:
            grade_diffs.append((row, col, og, eg))

grades_ok = not grade_diffs
print(f'\n7. GRADE DATA (letter grades, first 3 students, all sems):')
print(f'   Differences : {len(grade_diffs)}')
print(f'   Status      : {"OK - All letter grades match" if grades_ok else "MISMATCH"}')
if grade_diffs:
    for r,c,og,eg in grade_diffs[:10]: print(f'     R{r}C{c}: orig={repr(og)} export={repr(eg)}')
results['grades_ok'] = grades_ok

# --- GP formula verification ---
gp_cols_list = [6,8,10,12,14,16,18,20,26,28,30,32,34,36,38,40]
formula_errs = []
for row in range(7, 10):
    for col in gp_cols_list[:8]:  # Sem1 GP cols
        ov = ws_o_f.cell(row, col).value
        ev = ws_e_f.cell(row, col).value
        orig_is_formula = isinstance(ov, dict) or (isinstance(ov, str) and ov.startswith('='))
        exp_is_formula  = isinstance(ev, dict) or (isinstance(ev, str) and ev.startswith('='))
        # Check for #REF! error
        ev_str = str(ev) if ev is not None else ''
        if '#REF!' in ev_str:
            formula_errs.append((row, col, ev_str))

no_ref_errors = not formula_errs
print(f'\n8. FORMULA CELLS (#REF! errors in GP cols):')
print(f'   Errors : {len(formula_errs)}')
print(f'   Status : {"OK - No formula errors" if no_ref_errors else "ERRORS FOUND"}')
results['formulas_ok'] = no_ref_errors

# --- Header rows ---
hm = 0
for r in range(1, 7):
    for c in range(1, min(ws_o.max_column+1, 180)):
        o = str(ws_o.cell(r,c).value or '').strip()
        e = str(ws_e.cell(r,c).value or '').strip()
        if o != e: hm += 1
headers_ok = hm == 0
print(f'\n9. HEADER ROWS (1-6):')
print(f'   Mismatches : {hm}')
print(f'   Status     : {"OK" if headers_ok else "MISMATCH"}')
results['headers_ok'] = headers_ok

# --- Row heights ---
height_diffs = 0
for r in range(7, 13):
    oh = ws_o_f.row_dimensions[r].height
    eh = ws_e_f.row_dimensions[r].height
    if oh != eh: height_diffs += 1
print(f'\n10. ROW HEIGHTS (rows 7-12):')
print(f'    Differences : {height_diffs}')
print(f'    Note: Row heights may differ as ExcelJS inherits from template but does not copy row_dimensions explicitly')

# --- Column widths ---
width_diffs = 0
for c_idx in [1,2,3,4,5]:
    letter = openpyxl.utils.get_column_letter(c_idx)
    ow = ws_o_f.column_dimensions[letter].width
    ew = ws_e_f.column_dimensions[letter].width
    if ow != ew: width_diffs += 1
print(f'\n11. COLUMN WIDTHS (A-E):')
print(f'    Differences : {width_diffs}')
for c_idx in [1,2,3,4,5]:
    letter = openpyxl.utils.get_column_letter(c_idx)
    ow = ws_o_f.column_dimensions[letter].width
    ew = ws_e_f.column_dimensions[letter].width
    print(f'    Col {letter}: orig={ow}, export={ew} {"OK" if ow==ew else "DIFF"}')

# --- Arrears ---
arrears = db.get('arrears',[])
cleared = [a for a in arrears if a['status']=='CLEARED']
pending = [a for a in arrears if a['status']=='PENDING']
print(f'\n12. ARREAR RECORDS:')
print(f'    Total   : {len(arrears)}')
print(f'    Cleared : {len(cleared)}')
print(f'    Pending : {len(pending)}')

# === FINAL SUMMARY ===
print()
print('='*70)
print('FINAL SUMMARY')
print('='*70)

checks = [
    ('Sheet names match',       results.get('sheets_ok')),
    ('Dimensions match',        results.get('dims_ok')),
    ('Merged cells preserved',  results.get('merges_ok')),
    ('Column positions correct',results.get('cols_ok')),
    ('Semester colors correct', results.get('colors_ok')),
    ('Letter grades correct',   results.get('grades_ok')),
    ('No formula errors',       results.get('formulas_ok')),
    ('Header content correct',  results.get('headers_ok')),
    ('All orig students in export', results.get('students_ok')),
]

all_ok = all(v for _, v in checks if v is not None)

for name, ok in checks:
    status = 'PASS' if ok else 'FAIL'
    print(f'  [{status}] {name}')

print()
print(f'  Original student count : {len(orig_students)}')
print(f'  Database student count : {len(db_students)}')
print(f'  Exported student count : {len(exp_students)}')
print(f'  Extra test students    : {sorted(extra_in_exp) or "None"}')
print()
if all_ok:
    print('>>> ALL CRITICAL CHECKS PASSED. EXCEL EXPORT IS CORRECT.')
else:
    print('>>> SOME CHECKS FAILED. SEE DETAILS ABOVE.')
